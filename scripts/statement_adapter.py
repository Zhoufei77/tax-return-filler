#!/usr/bin/env python3
"""Read-only source recognition for Chinese profit-statement exports.

This adapter proposes cell mappings; it does not calculate or classify tax.
Source criteria, researched 2026-10-26:
* Qingdao Tax Bureau's 2025-01-23 FAQ distinguishes ASBE tax-filing
  '本期金额' (YTD) from Small ASBE tax-filing '本期金额' (current period):
  https://qingdao.chinatax.gov.cn/ssfg2019/cjwt/cjwtqt/202501/t20250123_92318.html
  That rule is NOT assumed for every accounting-software export.
* Kingdee Cloud Accounting V5 manual section 5.2 documents both current-month
  and YTD columns and Excel export. Chanjet permits column-order changes:
  https://cdn.jdy.com/jdy/helpDoc/ykj/云会计V5.0操作手册.pdf
  https://www.chanjetvip.com/doc/video/cate-4265/doc-36355.html
* Hubei Tax Bureau's public e-invoice manual, figure 199 / printed page 119,
  shows '发票基础信息' and '货物清单' as two sheets of an invoice export:
  https://www.hanyang.gov.cn/yxs/ztzl/qmszh/202305/P020230522408136179204.pdf

Headers and rows are matched by explicit aliases after whitespace, width and
Chinese ordinal normalization. No fuzzy matching, formula cache reading, zero
imputation, implicit period confirmation or account-number tax inference.
"""
from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import date, datetime
import json
from pathlib import Path
import re
import sys
import unicodedata

from tax_fill import Blocked, MAX_FILES, Source, amount, digest, dump_json, input_path, serial

VERSION = "0.1.0"
FIELD_ALIASES = {
    "operating_revenue": ("营业收入",),
    "operating_cost": ("营业成本",),
    "taxes_surcharges": ("税金及附加", "营业税金及附加"),
    "selling_expense": ("销售费用", "营业费用"),
    "administrative_expense": ("管理费用",),
    "rd_expense": ("研发费用",),
    "finance_expense": ("财务费用",),
    "other_income": ("其他收益",),
    "investment_income": ("投资收益",),
    "hedging_gain": ("净敞口套期收益",),
    "fair_value_gain": ("公允价值变动收益",),
    "credit_impairment": ("信用减值损失",),
    "asset_impairment": ("资产减值损失",),
    "asset_disposal": ("资产处置收益",),
    "operating_profit": ("营业利润",),
    "nonoperating_income": ("营业外收入",),
    "nonoperating_expense": ("营业外支出",),
    "profit_total": ("利润总额",),
}
ALIASES = {alias: field for field, names in FIELD_ALIASES.items() for alias in names}
YTD_HEADERS = {
    "本年累计", "本年累计金额", "本年累计数", "本年累计发生额",
    "年初至报告期末金额", "年初至报告期末数", "年初至报告期末累计金额",
    "年初至本期末金额",
}
MONTH_HEADERS = {"本月金额", "本月数", "本月发生额", "本月发生金额"}
QUARTER_HEADERS = {"本季金额", "本季数", "本季度金额", "本季度数", "本季度发生额"}
CONTEXT_HEADERS = {"本期金额", "本期数", "本期发生额", "本期发生金额"}
UNCLEAR_HEADERS = {"累计金额", "累计数", "本年金额", "本年数", "金额", "发生额"}
COMPARATIVE_HEADERS = {"上期金额", "上期数", "上年金额", "上年数", "上年同期金额", "上年同期累计金额"}
INVOICE_NUMBER_HEADERS = {"数电发票号码", "全电发票号码", "发票号码"}


def norm(value):
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(value)))


def semantic_label(value):
    """Return only whole main-row names; no contains/substring matching."""
    if not isinstance(value, str):
        return None
    text = norm(value)
    text = re.sub(r"^(?:[一二三四五六七八九十百零〇]+[、.．]|\([一二三四五六七八九十百零〇0-9]+\)|\d+[、.．])", "", text)
    if text.startswith("其中"):
        return None
    text = re.sub(r"^(?:加|减)[:：]", "", text)
    text = re.sub(r"\((?:亏损|净亏损|损失|净损失|净收益)[^()]{0,60}(?:填列|表示)\)$", "", text)
    return ALIASES.get(text)


def issue(code, message, **location):
    return {"code": code, "message": message, **location}


def header_meaning(value, standard=None, presentation=None):
    if not isinstance(value, str):
        return None
    text = norm(value)
    unit_match = re.search(r"\((万元|千元|元)\)$", text)
    unit = unit_match.group(1) if unit_match else None
    if unit_match:
        text = text[:unit_match.start()]
    if text in YTD_HEADERS:
        return {"basis": "ytd", "period_kind": None, "rule": "explicit_ytd_header", "unit": unit}
    if text in MONTH_HEADERS:
        return {"basis": "current", "period_kind": "month", "rule": "explicit_month_header", "unit": unit}
    if text in QUARTER_HEADERS:
        return {"basis": "current", "period_kind": "quarter", "rule": "explicit_quarter_header", "unit": unit}
    if text in CONTEXT_HEADERS:
        if presentation == "tax_filing" and standard in ("asbe", "small_asbe"):
            return {"basis": "ytd" if standard == "asbe" else "current", "period_kind": None,
                    "rule": "confirmed_tax_filing_and_accounting_standard", "unit": unit}
        return {"basis": "ambiguous", "period_kind": None,
                "rule": "current_header_requires_presentation_and_standard", "unit": unit}
    if text in UNCLEAR_HEADERS:
        return {"basis": "ambiguous", "period_kind": None, "rule": "header_does_not_define_period", "unit": unit}
    if text in COMPARATIVE_HEADERS:
        return {"basis": "comparative", "period_kind": None, "rule": "comparative_column_not_current", "unit": unit}
    return None


def parse_money(raw):
    conversions = []
    value = raw
    if isinstance(value, str):
        normalized = unicodedata.normalize("NFKC", value).strip()
        if normalized != value:
            conversions.append({"conversion": "width_and_edge_whitespace", "before": value, "after": normalized})
        value = normalized
        if value in ("", "-", "—", "–", "/"):
            raise Blocked("MISSING_AMOUNT", "空白或横线不是已确认零值")
        if re.fullmatch(r"\([0-9,]+(?:\.[0-9]+)?\)", value):
            converted = "-" + value[1:-1]
            conversions.append({"conversion": "accounting_parentheses_to_negative", "before": value, "after": converted})
            value = converted
    numeric = amount(value, "利润表金额")
    return format(numeric, "f"), conversions


def merged_columns(reader, sheet, cell, col):
    if not reader.workbook:
        return [col], None
    ws = reader.workbook[sheet]
    for merged in ws.merged_cells.ranges:
        if cell in merged:
            return list(range(merged.min_col, merged.max_col + 1)), str(merged)
    return [col], None


def selected_standard(cells, requested):
    evidence = []
    found = set()
    for entry in cells:
        value = entry["value"]
        if not isinstance(value, str):
            continue
        text = norm(value)
        if "小企业会计准则" in text or "会小企02" in text:
            found.add("small_asbe")
            evidence.append({"cell": entry["cell"], "text": value, "standard": "small_asbe"})
        remaining = text.replace("小企业会计准则", "")
        if "企业会计准则" in remaining:
            found.add("asbe")
            evidence.append({"cell": entry["cell"], "text": value, "standard": "asbe"})
    problems = []
    if len(found) > 1 or (requested and found and found != {requested}):
        problems.append(issue("ACCOUNTING_STANDARD_CONFLICT", "指定准则与表内准则冲突，或表内存在多个准则声明"))
        return None, evidence, problems
    return requested or (next(iter(found)) if found else None), evidence, problems


def metadata(cells):
    units = []
    candidates = []
    for entry in cells:
        value = entry["value"]
        if not isinstance(value, str):
            continue
        text = norm(value)
        match = re.search(r"(?<!编制)(?:金额单位|单位)[:：]?(?:人民币)?(万元|千元|元)(?:$|[;；,，])", text)
        if match:
            units.append({"cell": entry["cell"], "text": value, "unit": match.group(1)})
        if re.search(r"(?:20\d{2}年|20\d{2}[-/]\d{1,2}|编制单位|纳税人|统一社会信用代码|币种|会计期间|会计准则)", text):
            candidates.append({"cell": entry["cell"], "text": value, "confirmed": False})
    return units, candidates


def invoice_ids(reader, sheet):
    rows, cols = reader.dimensions(sheet)
    for row in range(1, min(rows, 15) + 1):
        for col in range(1, cols + 1):
            from openpyxl.utils.cell import get_column_letter
            address = f"{get_column_letter(col)}{row}"
            raw, kind = reader.raw(sheet, address)
            if isinstance(raw, str) and kind != "f" and norm(raw) in INVOICE_NUMBER_HEADERS:
                values = set()
                for rr in range(row + 1, rows + 1):
                    value, vk = reader.raw(sheet, f"{get_column_letter(col)}{rr}")
                    if isinstance(value, str) and value.strip() and vk != "f":
                        values.add(value.strip())
                return values
    return set()


def analyze_sheet(reader, sheet, path, sha, requested_standard=None, presentation=None):
    from openpyxl.utils.cell import get_column_letter
    rows, cols = reader.dimensions(sheet)
    all_cells, found_rows = [], []
    for row in range(1, rows + 1):
        for col in range(1, cols + 1):
            address = f"{get_column_letter(col)}{row}"
            raw, kind = reader.raw(sheet, address)
            if raw is None or raw == "":
                continue
            entry = {"cell": address, "row": row, "column": col, "value": serial(raw), "kind": kind}
            all_cells.append(entry)
            field = semantic_label(raw) if kind not in ("f", "e") else None
            if field:
                found_rows.append({"field": field, "row": row, "column": col, "cell": address, "label": raw})
    title = norm(sheet)
    classification = "invoice_base" if "发票基础信息" in title else "invoice_detail" if "货物清单" in title else "unknown"
    title_in_cells = any(isinstance(c["value"], str) and len(norm(c["value"])) < 100
                         and ("利润表" in norm(c["value"]) or "损益表" in norm(c["value"]))
                         and c["row"] <= min((r["row"] for r in found_rows), default=0)
                         for c in all_cells)
    is_profit = len({r["field"] for r in found_rows}) >= 2 or (found_rows and ("利润表" in title or "损益表" in title or title_in_cells))
    result = {"name": sheet, "classification": classification, "mappings": [], "issues": [], "warnings": []}
    if not is_profit:
        return result
    result["classification"] = "profit_statement"
    first_row = min(r["row"] for r in found_rows)
    prefix = [c for c in all_cells if c["row"] < first_row]
    standard, evidence, standard_problems = selected_standard(prefix, requested_standard)
    result.update(accounting_standard=standard, presentation=presentation, standard_evidence=evidence)
    result["issues"].extend(standard_problems)
    units, meta = metadata(prefix)
    result["metadata_candidates"] = meta
    result["unit_evidence"] = units
    unit_set = {u["unit"] for u in units}
    sheet_unit = next(iter(unit_set)) if len(unit_set) == 1 else None
    if len(unit_set) > 1:
        result["issues"].append(issue("UNIT_CONFLICT", "表内声明多个金额单位，需要确认每列单位"))
    headers = defaultdict(list)
    for entry in prefix:
        if entry["kind"] in ("f", "e"):
            continue
        meaning = header_meaning(entry["value"], standard, presentation)
        if not meaning:
            continue
        projected, merged = merged_columns(reader, sheet, entry["cell"], entry["column"])
        for col in projected:
            headers[col].append({"cell": entry["cell"], "row": entry["row"], "text": entry["value"],
                                 "column": col, "merged_range": merged, **meaning})
    selected, bad_columns = {}, set()
    for col, entries in headers.items():
        meanings = {(h["basis"], h["period_kind"]) for h in entries if h["basis"] not in ("ambiguous", "comparative")}
        chosen = max(entries, key=lambda h: h["row"])
        # A generic lower-level '金额' does not erase the explicit upper header.
        explicit = [h for h in entries if h["basis"] not in ("ambiguous", "comparative")]
        if explicit and chosen["basis"] == "ambiguous" and norm(chosen["text"]) in {"金额", "发生额"}:
            chosen = max(explicit, key=lambda h: h["row"])
        if len(meanings) > 1:
            bad_columns.add(col)
            result["issues"].append(issue("HEADER_CONFLICT", "同一列出现不同期间含义", column=get_column_letter(col)))
        chosen = dict(chosen)
        chosen["evidence"] = [{"cell": h["cell"], "text": h["text"]} for h in entries]
        selected[col] = chosen
    groups = defaultdict(list)
    for col, header in selected.items():
        if header["basis"] != "comparative":
            groups[(header["basis"], header["period_kind"], norm(header["text"]))].append(col)
    # Different explicit aliases for the same basis also need a deliberate choice.
    semantic_groups = defaultdict(list)
    for col, header in selected.items():
        if header["basis"] not in ("ambiguous", "comparative"):
            semantic_groups[(header["basis"], header["period_kind"])].append(col)
    duplicate_groups = [cols for cols in groups.values() if len(cols) > 1]
    duplicate_groups.extend(cols for cols in semantic_groups.values() if len(cols) > 1)
    seen_groups = set()
    for columns in duplicate_groups:
        key = tuple(sorted(columns))
        if key in seen_groups:
            continue
        seen_groups.add(key)
        bad_columns.update(columns)
        result["issues"].append(issue("DUPLICATE_HEADER", "多个列具有同一语义，不能自行任选一列", columns=[get_column_letter(c) for c in columns]))
    rows_by_field = defaultdict(list)
    for found in found_rows:
        rows_by_field[found["field"]].append(found)
    for field, entries in rows_by_field.items():
        if len(entries) > 1:
            result["issues"].append(issue("DUPLICATE_SEMANTIC_ROW", "同一财务主行出现多次，不能自动选择或相加", field=field, cells=[r["cell"] for r in entries]))
    if not selected:
        result["issues"].append(issue("MISSING_PERIOD_HEADER", "未找到明确期间金额列，不要求用户改表，需确认列含义"))
    for found in found_rows:
        for col, header in selected.items():
            if header["basis"] == "comparative" or col <= found["column"]:
                continue
            address = f"{get_column_letter(col)}{found['row']}"
            mapping = {"field": found["field"], "basis": header["basis"], "period_kind": header["period_kind"],
                       "source": {"path": str(path), "sha256": sha, "sheet": sheet, "cell": address, "label_cell": found["cell"]},
                       "label": found["label"], "header": header, "raw_value": None, "raw_numeric": None,
                       "conversions": [], "amount_unit": header["unit"] or sheet_unit, "usable": True, "issues": []}
            if len(rows_by_field[found["field"]]) > 1:
                mapping["issues"].append(issue("DUPLICATE_SEMANTIC_ROW", "财务主行重复"))
            if col in bad_columns:
                mapping["issues"].append(issue("AMBIGUOUS_HEADER", "期间列有冲突或重名"))
            if header["basis"] == "ambiguous":
                mapping["issues"].append(issue("AMBIGUOUS_PERIOD_BASIS", "列名本身不足以确定本期或累计；请一次确认报表来源语境及会计准则"))
            if standard_problems:
                mapping["issues"].append(issue("ACCOUNTING_STANDARD_CONFLICT", "会计准则声明冲突"))
            if mapping["amount_unit"] is None:
                mapping["issues"].append(issue("UNIT_UNCONFIRMED", "未识别明确金额单位，不能默认元"))
            try:
                raw, _ = reader.raw(sheet, address)
                mapping["raw_value"] = serial(raw)
                raw = reader.value(sheet, address)
                mapping["raw_numeric"], mapping["conversions"] = parse_money(raw)
            except Blocked as error:
                mapping["issues"].append(issue(error.code, error.message))
            mapping["usable"] = not mapping["issues"]
            result["mappings"].append(mapping)
    for mapping in result["mappings"]:
        for problem in mapping["issues"]:
            result["issues"].append({**problem, "field": mapping["field"], "cell": mapping["source"]["cell"], "header_cell": mapping["header"]["cell"]})
    return result


def analyze_file(path, requested_standard=None, presentation=None):
    sha = digest(path)
    result = {"path": str(path), "sha256": sha, "sheets": [], "issues": [], "warnings": []}
    reader = None
    try:
        reader = Source(path)
        for sheet in reader.sheets:
            result["sheets"].append(analyze_sheet(reader, sheet, path, sha, requested_standard, presentation))
        base = [s for s in result["sheets"] if s["classification"] == "invoice_base"]
        detail = [s for s in result["sheets"] if s["classification"] == "invoice_detail"]
        if base and detail:
            common = invoice_ids(reader, base[0]["name"]) & invoice_ids(reader, detail[0]["name"])
            result["warnings"].append(issue("INVOICE_DUAL_GRAIN", "发票基础信息与货物清单是票级/行级两种粒度，不能将两张表金额相加；不在此自动税务分类",
                                            sheets=[base[0]["name"], detail[0]["name"]], matching_invoice_ids_count=len(common),
                                            overlap_evidence="matching_text_invoice_ids" if common else "sheet_structure_only"))
    except Blocked as error:
        result["issues"].append(issue(error.code, error.message))
    except Exception as error:
        result["issues"].append(issue("ADAPTER_INPUT_ERROR", f"{type(error).__name__}: {error}"))
    finally:
        if reader and reader.workbook:
            reader.workbook.close()
    return result


def scan_directory(directory, accounting_standard=None, presentation=None):
    directory = Path(directory).expanduser().resolve()
    if not directory.is_dir():
        raise Blocked("INPUT_DIRECTORY", "输入必须是本次资料目录")
    candidates = sorted(p for p in directory.rglob("*") if p.is_file() and p.suffix.lower() in (".xlsx", ".csv"))
    if len(candidates) > MAX_FILES:
        raise Blocked("FILE_LIMIT", "资料文件数超过安全扫描上限")
    report = {"schema_version": 1, "adapter_version": VERSION, "root": str(directory),
              "accounting_standard": accounting_standard, "presentation": presentation,
              "files": [], "source_mappings": [], "issues": [], "warnings": [],
              "tax_calculation_performed": False, "inputs_changed": False}
    seen = {}
    for raw in candidates:
        try:
            path = input_path(directory, str(raw))
            sha = digest(path)
            if sha in seen:
                report["issues"].append(issue("DUPLICATE_FILE_CONTENT", "重复文件内容不重复识别或相加", path=str(path), first_path=seen[sha]))
                continue
            seen[sha] = str(path)
            file_result = analyze_file(path, accounting_standard, presentation)
            report["files"].append(file_result)
            report["issues"].extend({**p, "path": str(path)} for p in file_result["issues"])
            report["warnings"].extend({**p, "path": str(path)} for p in file_result["warnings"])
            for sheet in file_result["sheets"]:
                report["source_mappings"].extend(sheet["mappings"])
                report["issues"].extend({**p, "path": str(path), "sheet": sheet["name"]} for p in sheet["issues"])
            if digest(path) != sha:
                report["inputs_changed"] = True
                report["issues"].append(issue("HASH_CHANGED", "扫描期间输入变化，请重扫并复核", path=str(path)))
        except Blocked as error:
            report["issues"].append(issue(error.code, error.message, path=str(raw)))
    if not report["source_mappings"]:
        report["issues"].append(issue("NO_PROFIT_STATEMENT", "没有识别到可建立映射的利润表；无需用户改固定表头，请检查资料类型或补充来源语境"))
    report["usable_mapping_count"] = sum(m["usable"] for m in report["source_mappings"])
    report["status"] = "NEEDS_CONFIRMATION" if report["issues"] else "MAPPINGS_PROPOSED_FOR_REVIEW"
    return report


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    sp = sub.add_parser("scan", help="只读识别利润表，输出坐标、金额、来源及歧义")
    sp.add_argument("input_dir", type=Path)
    sp.add_argument("--out", type=Path, required=True)
    sp.add_argument("--accounting-standard", choices=("asbe", "small_asbe"))
    sp.add_argument("--presentation", choices=("tax_filing", "software_export"))
    args = parser.parse_args(argv)
    try:
        out = args.out.expanduser().absolute()
        if out.exists() or out.is_symlink():
            raise Blocked("OUTPUT_EXISTS", "--out 必须为不存在的新 JSON 文件，不允许覆盖")
        report = scan_directory(args.input_dir, args.accounting_standard, args.presentation)
        out.parent.mkdir(parents=True, exist_ok=True)
        dump_json(out, report)
        print(json.dumps({"status": report["status"], "files": len(report["files"]),
                          "usable_mappings": report["usable_mapping_count"], "issues": len(report["issues"]),
                          "output": str(out)}, ensure_ascii=False))
        return 2 if report["issues"] else 0
    except Blocked as error:
        print(json.dumps({"status": "BLOCKED", "code": error.code, "message": error.message}, ensure_ascii=False), file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
