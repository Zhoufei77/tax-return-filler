#!/usr/bin/env python3
"""Deterministic, offline, review-required return-template filler.

Python >= 3.10; openpyxl is the only third-party dependency.
Money expressions operate in CNY yuan; money targets are converted to the
job context amount_unit and rounded HALF_UP to 2 decimals only at the boundary.
No tax rules, OCR, formula evaluation, network access or filing are implemented.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import html
import io
import json
import re
import sys
import zipfile
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP, localcontext
from pathlib import Path

VERSION = "0.1.0"
STATUS = "DRAFT_REVIEW_REQUIRED"
UNITS = {"元": Decimal(1), "千元": Decimal(1000), "万元": Decimal(10000)}
MAX_FILE = 25 * 1024 * 1024
MAX_EXPANDED = 150 * 1024 * 1024
MAX_CELLS = 200000
MAX_ROWS = 50000
MAX_FILES = 100
MAX_EXPR_NODES = 5000
MAX_AMOUNT = Decimal("1e20")
CELL_RE = re.compile(r"^[A-Z]{1,3}[1-9][0-9]{0,6}$")
NUMBER_RE = re.compile(r"^[+-]?(?:[0-9]+|[0-9]{1,3}(?:,[0-9]{3})+)(?:\.[0-9]+)?$")
SAFE_FORMULA_FUNCTIONS = frozenset("""
SUM SUMIF SUMIFS SUMPRODUCT IF IFS IFERROR IFNA AND OR NOT XOR TRUE FALSE
ROUND ROUNDUP ROUNDDOWN MROUND ABS MIN MAX INT MOD TRUNC SIGN PRODUCT POWER
SQRT CEILING CEILING.MATH CEILING.PRECISE FLOOR FLOOR.MATH FLOOR.PRECISE
COUNT COUNTA COUNTBLANK COUNTIF COUNTIFS AVERAGE AVERAGEIF AVERAGEIFS
SUBTOTAL AGGREGATE MEDIAN LARGE SMALL RANK RANK.EQ RANK.AVG
DATE YEAR MONTH DAY DAYS EDATE EOMONTH DATEVALUE WEEKDAY WEEKNUM TODAY
TEXT VALUE NUMBERVALUE LEFT RIGHT MID LEN TRIM CLEAN UPPER LOWER PROPER
CONCAT CONCATENATE TEXTJOIN REPLACE SUBSTITUTE EXACT SEARCH FIND REPT
ISBLANK ISNUMBER ISTEXT ISERROR ISERR ISNA ISLOGICAL ISEVEN ISODD
INDEX MATCH XMATCH XLOOKUP VLOOKUP HLOOKUP LOOKUP CHOOSE ROW ROWS COLUMN COLUMNS
""".split())


def safe_formula(value, location):
    """Allow only inert references/operators and known pure Excel functions."""
    require(isinstance(value, str), "TEMPLATE_FORMULA_UNSUPPORTED", f"拒绝数组/数据表等未知公式形式: {location}")
    require(len(value) <= 32767, "TEMPLATE_FORMULA_UNSUPPORTED", f"公式过长: {location}")
    require(not re.search(r"[\[\]\\|]|[A-Za-z][A-Za-z0-9+.-]*://|(?:[\"'])\s*[A-Za-z][A-Za-z0-9+.-]*:|[A-Za-z]:[/\\]", value),
            "TEMPLATE_UNSAFE_FORMULA", f"模板公式含外部路径、URI、外部引用或 DDE: {location}")
    require(not re.search(r"[\"']\s*(?:/|\.\.?[/\\])", value), "TEMPLATE_UNSAFE_FORMULA", f"模板公式含外部文件路径: {location}")
    from openpyxl.formula.tokenizer import Tokenizer
    try:
        tokens = Tokenizer(value if value.startswith("=") else "=" + value).items
    except Exception:
        fail("TEMPLATE_FORMULA_UNSUPPORTED", f"无法安全解析模板公式: {location}")
    for token in tokens:
        if token.type == "FUNC" and token.subtype == "OPEN":
            function = token.value[:-1].strip().upper()
            require(function in SAFE_FORMULA_FUNCTIONS, "TEMPLATE_UNSAFE_FORMULA", f"模板函数不在纯计算白名单: {function} ({location})")


def validate_template_formulas(workbook):
    for ws in workbook:
        for row in ws:
            for cell in row:
                if cell.data_type == "f":
                    safe_formula(cell.value, f"{ws.title}!{cell.coordinate}")
        for name in ws.defined_names.values():
            safe_formula(name.attr_text, f"{ws.title} 定义名称 {name.name}")
    for name in workbook.defined_names.values():
        safe_formula(name.attr_text, f"工作簿定义名称 {name.name}")


class Blocked(Exception):
    def __init__(self, code, message):
        self.code, self.message = code, str(message)
        super().__init__(self.message)


def fail(code, message):
    raise Blocked(code, message)


def require(condition, code, message):
    if not condition:
        fail(code, message)


def digest(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def cell_address(value):
    require(isinstance(value, str), "CELL_ADDRESS", "单元格地址必须是字符串")
    value = value.upper()
    require(bool(CELL_RE.fullmatch(value)), "CELL_ADDRESS", f"无效单元格地址: {value}")
    from openpyxl.utils.cell import coordinate_to_tuple
    row, col = coordinate_to_tuple(value)
    require(row <= 1048576 and col <= 16384, "CELL_ADDRESS", f"单元格越界: {value}")
    return value


def iso_date(value, label):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    require(isinstance(value, str), "DATE_VALUE", f"{label} 必须为 ISO 日期 YYYY-MM-DD")
    require(bool(re.fullmatch(r"\d{4}-\d{2}-\d{2}", value)), "DATE_VALUE", f"{label} 必须为 ISO 日期 YYYY-MM-DD")
    try:
        return date.fromisoformat(value)
    except ValueError:
        fail("DATE_VALUE", f"{label} 不是有效日期: {value}")


def amount(value, label):
    require(value is not None and value != "", "MISSING_VALUE", f"{label} 缺少数值，不能补零")
    require(not isinstance(value, bool) and not isinstance(value, (date, datetime)), "NUMBER_TYPE", f"{label} 不是金额")
    if isinstance(value, str):
        value = value.strip()
        require(bool(NUMBER_RE.fullmatch(value)), "NUMBER_FORMAT", f"{label} 不是明确数值: {value}")
        value = value.replace(",", "")
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError):
        fail("NUMBER_FORMAT", f"{label} 不是有效数值")
    require(result.is_finite() and abs(result) <= MAX_AMOUNT, "NUMBER_RANGE", f"{label} 数值非有限或超限")
    return result


def check_unit(unit):
    require(unit in UNITS, "AMOUNT_UNIT", "金额单位必须明确为 元、千元 或 万元")
    return UNITS[unit]


def serial(value):
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def context_text(value, role):
    if role in ("period_start", "period_end") and isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(serial(value))


def dump_json(path, value):
    with path.open("x", encoding="utf-8") as f:
        json.dump(value, f, ensure_ascii=False, indent=2, default=serial, allow_nan=False)
        f.write("\n")


def input_path(base, value):
    require(isinstance(value, str) and value.strip(), "INPUT_PATH", "缺少文件路径")
    raw = Path(value).expanduser()
    if not raw.is_absolute():
        raw = base / raw
    require(not raw.is_symlink(), "SYMLINK_INPUT", f"拒绝符号链接输入: {raw}")
    path = raw.resolve()
    require(path.is_file(), "INPUT_MISSING", f"文件不存在: {path}")
    require(path.stat().st_size <= MAX_FILE, "FILE_LIMIT", f"文件超过 {MAX_FILE} 字节: {path.name}")
    return path


def verify_hash(path, expected):
    require(isinstance(expected, str) and re.fullmatch(r"[a-fA-F0-9]{64}", expected), "HASH_REQUIRED", f"必须提供文件 SHA256: {path.name}")
    actual = digest(path)
    require(actual == expected.lower(), "HASH_CHANGED", f"文件 SHA256 不匹配，必须重新扫描复核: {path.name}")
    return actual


def inspect_xlsx(path, template=False):
    require(path.suffix.lower() == ".xlsx", "XLSX_ONLY", "工作簿仅支持 .xlsx；不接受宏、旧版 XLS 或改名文件")
    require(zipfile.is_zipfile(path), "XLSX_INVALID", f"不是有效 XLSX: {path.name}")
    try:
        with zipfile.ZipFile(path) as z:
            infos = z.infolist()
            require(len(infos) <= 5000, "ZIP_LIMIT", "工作簿 ZIP 成员过多")
            require(sum(i.file_size for i in infos) <= MAX_EXPANDED, "ZIP_LIMIT", "工作簿解压大小超限")
            names = [i.filename.lower() for i in infos]
            require(len(set(names)) == len(names), "ZIP_DUPLICATE", "ZIP 含重复或大小写冲突的成员名")
            require({"[content_types].xml", "_rels/.rels", "xl/workbook.xml"} <= set(names), "XLSX_INVALID", "缺少必要 OOXML 成员")
            xml_cell_count = 0
            for i in infos:
                require(not (i.flag_bits & 1), "ENCRYPTED_FILE", "不支持加密 ZIP/XLSX")
                require(i.file_size <= MAX_EXPANDED and i.file_size / max(i.compress_size, 1) <= 2000,
                        "ZIP_LIMIT", "疑似异常压缩工作簿")
                require(not i.filename.startswith(("/", "\\")) and not re.match(r"^[A-Za-z]:", i.filename)
                        and "\\" not in i.filename and ".." not in Path(i.filename).parts,
                        "ZIP_PATH", "ZIP 成员路径不安全")
                if i.filename.lower().endswith((".xml", ".rels")):
                    tail = b""
                    with z.open(i) as xml:
                        for chunk in iter(lambda: xml.read(65536), b""):
                            require(b"\x00" not in chunk, "XML_ENCODING", "仅接受不含 NUL 的 UTF-8 OOXML")
                            token = (tail + chunk).upper()
                            require(b"<!DOCTYPE" not in token and b"<!ENTITY" not in token,
                                    "XML_ENTITY", "拒绝包含 DTD 或实体声明的 OOXML")
                            if i.filename.lower().startswith("xl/worksheets/"):
                                # Bound actual XML cell elements before openpyxl materializes them.
                                complete = tail + chunk
                                xml_cell_count += sum(m.end() > len(tail) for m in re.finditer(
                                    rb"<(?:[A-Za-z_][A-Za-z_0-9]*:)?c(?:\s|>)", complete))
                                require(xml_cell_count <= MAX_CELLS, "SHEET_LIMIT", "工作簿 XML 单元格数超限")
                            tail = (tail + chunk)[-64:]
            require(not any("vbaproject" in n or n.startswith("xl/macrosheets/") for n in names), "MACRO_FILE", "拒绝含宏工作簿")
            require(not any(n.startswith("xl/externallinks/") for n in names), "EXTERNAL_LINK", "拒绝含外部工作簿链接的文件")
            if template:
                require(not any(n.startswith(("xl/drawings/", "xl/charts/", "xl/media/", "xl/activex/", "xl/embeddings/"))
                                or "vml" in n or "threadedcomments" in n for n in names),
                        "TEMPLATE_OBJECTS", "模板含形状、图表、图片、控件或其他不可安全保留的对象")
                require(not any(n.startswith(("_xmlsignatures/", "xl/connections", "xl/querytables/", "xl/pivottables/", "xl/pivotcache/")) for n in names),
                        "TEMPLATE_OBJECTS", "模板含数字签名、查询连接或透视表，无法保证保存后完整")
    except (zipfile.BadZipFile, OSError) as e:
        fail("XLSX_INVALID", f"工作簿 ZIP 无效: {e}")


class Source:
    def __init__(self, path, template=False):
        self.path, self.workbook = path, None
        if path.suffix.lower() == ".csv" and not template:
            data = path.read_bytes()
            try:
                text = data.decode("utf-8-sig")
                self.encoding = "utf-8-sig"
            except UnicodeDecodeError:
                try:
                    text = data.decode("gb18030")
                    self.encoding = "gb18030"
                except UnicodeDecodeError:
                    fail("CSV_ENCODING", "CSV 必须使用 UTF-8 或 GB18030 编码")
            require("\x00" not in text, "CSV_INVALID", "CSV 含 NUL 字符")
            self.rows, count = [], 0
            for row in csv.reader(io.StringIO(text)):
                count += len(row)
                require(len(self.rows) < MAX_ROWS and count <= MAX_CELLS, "SHEET_LIMIT", "CSV 行数或单元格数超限")
                self.rows.append(row)
            self.sheets = ["CSV"]
        else:
            inspect_xlsx(path, template)
            from openpyxl import load_workbook
            try:
                self.workbook = load_workbook(path, data_only=False, keep_links=False)
            except Exception as e:
                fail("XLSX_LOAD", f"无法读取 XLSX: {e}")
            self.sheets = self.workbook.sheetnames
            total = 0
            for ws in self.workbook:
                require(ws.max_row <= MAX_ROWS and ws.max_row * ws.max_column <= MAX_CELLS,
                        "SHEET_LIMIT", f"工作表范围过大: {ws.title}")
                total += ws.max_row * ws.max_column
            require(total <= MAX_CELLS, "SHEET_LIMIT", "工作簿单元格范围超限")
            if template:
                security = self.workbook.security
                require(security is None or (not security.lockStructure and not security.lockWindows),
                        "PROTECTED_TEMPLATE", "拒绝受保护工作簿")
                require(not any(ws.protection.sheet for ws in self.workbook), "PROTECTED_TEMPLATE", "拒绝含受保护工作表的模板")
                validate_template_formulas(self.workbook)

    def dimensions(self, sheet):
        require(sheet in self.sheets, "SHEET_MISSING", f"工作表不存在: {sheet}")
        if self.workbook:
            ws = self.workbook[sheet]
            return ws.max_row, ws.max_column
        return len(self.rows), max(map(len, self.rows), default=0)

    def raw(self, sheet, address):
        address = cell_address(address)
        self.dimensions(sheet)
        if self.workbook:
            c = self.workbook[sheet][address]
            return c.value, c.data_type
        from openpyxl.utils.cell import coordinate_to_tuple
        row, col = coordinate_to_tuple(address)
        value = self.rows[row - 1][col - 1] if row <= len(self.rows) and col <= len(self.rows[row - 1]) else None
        kind = "f" if isinstance(value, str) and value.startswith("=") else "s"
        return value, kind

    def value(self, sheet, address):
        value, kind = self.raw(sheet, address)
        require(kind != "f", "FORMULA_SOURCE", f"源单元格含公式，不读取缓存或猜算: {sheet}!{address}")
        require(kind != "e", "CELL_ERROR", f"源单元格为 Excel 错误: {sheet}!{address}")
        require(value is not None and (not isinstance(value, str) or value.strip() != ""),
                "MISSING_VALUE", f"源单元格为空: {sheet}!{address}")
        return value


def scan(directory, output):
    directory = directory.expanduser().resolve()
    require(directory.is_dir(), "INPUT_PATH", "scan 输入必须为资料目录")
    require(not output.exists(), "OUTPUT_EXISTS", "扫描输出文件已存在，不允许覆盖")
    candidates = sorted(p for p in directory.rglob("*") if p.is_file() and p.suffix.lower() in (".xlsx", ".csv"))
    require(len(candidates) <= MAX_FILES, "FILE_LIMIT", "扫描资料文件数超限")
    result = {"schema_version": 1, "engine_version": VERSION, "root": str(directory), "files": [], "issues": []}
    for raw in candidates:
        try:
            path = input_path(directory, str(raw))
            require(path.resolve() != output.resolve(), "OUTPUT_INPUT", "输出不能覆盖扫描输入")
            source = Source(path)
            item = {"path": str(path.relative_to(directory)), "sha256": digest(path), "size": path.stat().st_size, "sheets": []}
            for sheet in source.sheets:
                rows, cols = source.dimensions(sheet)
                cells = []
                from openpyxl.utils.cell import get_column_letter
                for row in range(1, rows + 1):
                    for col in range(1, cols + 1):
                        address = f"{get_column_letter(col)}{row}"
                        value, kind = source.raw(sheet, address)
                        if value is not None and value != "":
                            cells.append({"cell": address, "value": serial(value), "type": kind,
                                          "usable": kind not in ("f", "e"), "note": "公式不可取数" if kind == "f" else ""})
                item["sheets"].append({"name": sheet, "max_row": rows, "max_column": cols, "cells": cells})
            result["files"].append(item)
            if source.workbook:
                source.workbook.close()
        except Blocked as e:
            result["issues"].append({"file": str(raw), "code": e.code, "message": e.message})
        except Exception as e:
            result["issues"].append({"file": str(raw), "code": "SCAN_FAILED", "message": str(e)})
    output.parent.mkdir(parents=True, exist_ok=True)
    dump_json(output, result)
    return 2 if result["issues"] else 0


class Engine:
    def __init__(self, job, base, audit):
        self.job, self.base, self.audit = job, base, audit
        self.sources, self.fields, self.traces = {}, {}, []
        self.field_records = {}
        self.context_bindings = {}
        self.nodes = 0
        self.context = job.get("context", {})

    def validate(self):
        require(self.job.get("schema_version") == 1, "SCHEMA_VERSION", "schema_version 必须为 1")
        c = self.context
        for key in ("taxpayer_id", "period_start", "period_end", "period_kind", "form_name", "form_version", "region", "currency", "amount_unit"):
            require(isinstance(c.get(key), str) and c[key].strip(), "CONTEXT_MISSING", f"缺少明确上下文: {key}")
        require(c["currency"] == "CNY", "CURRENCY", "仅支持 CNY，不自动换汇")
        check_unit(c["amount_unit"])
        self.start, self.end = iso_date(c["period_start"], "period_start"), iso_date(c["period_end"], "period_end")
        require(self.start <= self.end, "PERIOD", "所属期起止倒置")
        require(c["period_kind"] in ("month", "quarter", "year"), "PERIOD_KIND", "period_kind 必须为 month、quarter 或 year")
        from calendar import monthrange
        require(self.start.day == 1 and self.end.day == monthrange(self.end.year, self.end.month)[1], "PERIOD", "所属期必须覆盖完整月/季/年")
        if c["period_kind"] == "month":
            require((self.start.year, self.start.month) == (self.end.year, self.end.month), "PERIOD", "month 必须为同一完整月")
        elif c["period_kind"] == "quarter":
            require(self.start.year == self.end.year and self.start.month in (1, 4, 7, 10) and self.end.month == self.start.month + 2,
                    "PERIOD", "quarter 必须为一个自然季度")
        else:
            require(self.start.year == self.end.year and self.start.month == 1 and self.end.month == 12, "PERIOD", "year 必须为一个自然年")
        require(self.end <= date.today(), "FUTURE_PERIOD", "不生成尚未结束期间的完整申报草稿")
        confirmations = self.job.get("confirmations", {})
        require(confirmations.get("context_confirmed") is True and confirmations.get("materials_complete") is True,
                "HUMAN_CONFIRMATION", "须由用户明确确认主体/所属期/表单版本及材料完整性，不能由 Agent 代为假定")
        require(isinstance(self.job.get("sources"), list) and self.job["sources"], "SOURCES_MISSING", "至少需要一个来源")
        source_paths, source_hashes = set(), set()
        for meta in self.job["sources"]:
            sid = meta.get("id")
            require(isinstance(sid, str) and sid and sid not in self.sources, "SOURCE_ID", "来源 id 缺失或重复")
            sc = meta.get("context", {})
            require(sc.get("taxpayer_id") == c["taxpayer_id"], "SOURCE_TAXPAYER", f"来源主体不一致: {sid}")
            require(sc.get("currency") == "CNY", "SOURCE_CURRENCY", f"来源币种不明确或不一致: {sid}")
            check_unit(sc.get("amount_unit"))
            ss, se = iso_date(sc.get("period_start"), f"{sid} period_start"), iso_date(sc.get("period_end"), f"{sid} period_end")
            require(ss <= se and ss <= self.start and se >= self.end, "SOURCE_PERIOD", f"来源数据窗口未覆盖目标期间: {sid}")
            anchors = meta.get("identity_anchors", [])
            require(isinstance(anchors, list), "IDENTITY_ANCHOR", "identity_anchors 必须为列表")
            require(meta.get("metadata_confirmed") is True, "SOURCE_CONFIRMATION", f"用户尚未确认来源主体、期间与单位: {sid}")
            path = input_path(self.base, meta.get("path"))
            actual = verify_hash(path, meta.get("sha256"))
            require(path not in source_paths and actual not in source_hashes, "SOURCE_DUPLICATE", "同一路径或相同内容不能以多个来源 id 重复计入")
            source_paths.add(path)
            source_hashes.add(actual)
            source = Source(path)
            for anchor in anchors:
                require("expected" in anchor, "IDENTITY_ANCHOR", "身份锚点缺少 expected")
                actual_value = source.value(anchor.get("sheet"), cell_address(anchor.get("cell")))
                require(str(serial(actual_value)) == str(anchor["expected"]), "IDENTITY_MISMATCH", f"来源身份锚点不匹配: {sid}")
            self.sources[sid] = {"meta": meta, "reader": source, "start": ss, "end": se}
            self.audit["inputs"].append({"id": sid, "path": str(path), "sha256": actual, "context": sc,
                                         "identity_evidence": "anchors_and_human_declaration" if anchors else "human_declaration_only"})
        tm = self.job.get("template", {})
        self.template_path = input_path(self.base, tm.get("path"))
        actual = verify_hash(self.template_path, tm.get("sha256"))
        self.template = Source(self.template_path, template=True)
        self.audit["inputs"].append({"id": "template", "path": str(self.template_path), "sha256": actual})
        require(self.template_path not in [s["reader"].path for s in self.sources.values()], "TEMPLATE_SOURCE", "模板不能同时用作取数来源")
        bindings = tm.get("context_bindings")
        require(isinstance(bindings, dict) and "taxpayer_id" in bindings,
                "TEMPLATE_CONTEXT_BINDINGS", "模板必须绑定 taxpayer_id 及期间")
        combined_period = "period" in bindings
        if combined_period:
            require("period_start" not in bindings and "period_end" not in bindings,
                    "TEMPLATE_CONTEXT_BINDINGS", "期间绑定须二选一：period 单格或 period_start/period_end 两格")
        else:
            require({"period_start", "period_end"} <= set(bindings), "TEMPLATE_CONTEXT_BINDINGS", "分格期间必须同时绑定 period_start 和 period_end")
        seen_bindings = set()
        self.audit["template_context_bindings"] = []
        for role, binding in bindings.items():
            require((role in c or role == "period") and isinstance(binding, dict), "TEMPLATE_CONTEXT_BINDINGS", f"无效模板上下文绑定: {role}")
            require("expected" not in binding, "TEMPLATE_CONTEXT_BINDINGS", "上下文绑定不允许自由 expected，必须由任务上下文生成")
            if role == "period":
                period_format = binding.get("format", "iso_range")
                require(period_format in ("iso_range", "cn_range"), "TEMPLATE_PERIOD_FORMAT", "period.format 只能为 iso_range 或 cn_range")
                if period_format == "iso_range":
                    expected = f"{c['period_start']} 至 {c['period_end']}"
                else:
                    expected = f"{self.start.year}年{self.start.month}月{self.start.day}日至{self.end.year}年{self.end.month}月{self.end.day}日"
            else:
                require("format" not in binding, "TEMPLATE_PERIOD_FORMAT", "format 仅适用于 period 单格绑定")
                expected = c[role]
            sheet, cell = binding.get("sheet"), cell_address(binding.get("cell"))
            key = (sheet, cell)
            require(key not in seen_bindings, "TEMPLATE_CONTEXT_BINDINGS", "不同身份/期间角色不能绑定同一目标格")
            seen_bindings.add(key)
            raw, kind = self.template.raw(sheet, cell)
            require(kind not in ("f", "e"), "TEMPLATE_CONTEXT_MISMATCH", f"模板身份/期间格不能为公式或错误值: {role}")
            from openpyxl.cell.cell import MergedCell
            require(not isinstance(self.template.workbook[sheet][cell], MergedCell), "TEMPLATE_CONTEXT_BINDINGS", "模板身份/期间格不能是合并区域非首格")
            blank = raw is None or raw == ""
            if not blank:
                require((isinstance(raw, str) or (role in ("period_start", "period_end") and isinstance(raw, (date, datetime))))
                        and context_text(raw, role) == expected,
                        "TEMPLATE_CONTEXT_MISMATCH", f"模板原有身份/期间与任务不一致: {role} ({sheet}!{cell})")
            self.context_bindings[key] = {"role": role, "expected": expected, "blank": blank, "existing": serial(raw)}
            self.audit["template_context_bindings"].append({"role": role, "sheet": sheet, "cell": cell,
                "expected": expected, "existing": serial(raw), "requires_text_field": blank,
                "format": binding.get("format", "iso_range") if role == "period" else None})

    def period_for(self, field):
        basis = field.get("basis", "current")
        require(basis in ("current", "ytd"), "FIELD_BASIS", "basis 只能为 current 或 ytd")
        return (self.start if basis == "current" else date(self.start.year, 1, 1)), self.end

    def source(self, sid):
        require(sid in self.sources, "SOURCE_MISSING", f"来源不存在: {sid}")
        return self.sources[sid]

    def expression(self, expr, field, depth=0, record_keys=None):
        if record_keys is None:
            record_keys = set()
        require(depth <= 30, "EXPRESSION_LIMIT", "表达式嵌套过深")
        self.nodes += 1
        require(self.nodes <= MAX_EXPR_NODES, "EXPRESSION_LIMIT", "表达式节点过多")
        require(isinstance(expr, dict), "EXPRESSION", "表达式必须为对象")
        op = expr.get("op")
        if op == "ref":
            s = self.source(expr.get("source"))
            start, end = self.period_for(field)
            scope_evidence = None
            if "period_scope" in expr:
                scope = expr["period_scope"]
                require(isinstance(scope, dict), "REF_PERIOD_SCOPE", "ref.period_scope 必须为包含起止期和期间表头锚点的对象")
                scoped_start = iso_date(scope.get("period_start"), "ref.period_scope.period_start")
                scoped_end = iso_date(scope.get("period_end"), "ref.period_scope.period_end")
                require((scoped_start, scoped_end) == (start, end), "REF_PERIOD_SCOPE", "ref.period_scope 必须严格匹配字段 current/ytd 所需期间")
                require(s["start"] <= start and s["end"] >= end, "REF_PERIOD_SCOPE", "来源数据窗口未覆盖 ref.period_scope")
                anchor = scope.get("anchor", {})
                require(isinstance(anchor, dict) and isinstance(anchor.get("text"), str) and anchor["text"],
                        "REF_PERIOD_ANCHOR", "ref.period_scope 需要真实期间列表头 anchor.cell 与 anchor.text")
                anchor_cell = cell_address(anchor.get("cell"))
                anchor_value = s["reader"].value(expr.get("sheet"), anchor_cell)
                require(anchor_value == anchor["text"], "REF_PERIOD_ANCHOR", "ref 期间列表头锚点不匹配")
                scope_evidence = {"period_start": scoped_start.isoformat(), "period_end": scoped_end.isoformat(),
                                  "anchor": {"sheet": expr.get("sheet"), "cell": anchor_cell, "text": anchor["text"], "observed": serial(anchor_value)}}
            else:
                require(s["start"] == start and s["end"] == end, "REF_PERIOD", "ref 来源期间必须与字段 current/ytd 完全一致；同文件混合期间列需提供 period_scope 与真实表头锚点")
            raw = s["reader"].value(expr.get("sheet"), cell_address(expr.get("cell")))
            unit = expr.get("amount_unit", s["meta"]["context"]["amount_unit"])
            result = amount(raw, "ref") * check_unit(unit)
            self.traces.append({"op": op, "source": expr["source"], "sheet": expr["sheet"], "cell": expr["cell"],
                                "raw": serial(raw), "amount_unit": unit, "result_yuan": result, "period_scope": scope_evidence})
        elif op == "const":
            require(isinstance(expr.get("reason"), str) and expr["reason"].strip(), "CONST_REASON", "常量必须附来源/用途 reason")
            result = amount(expr.get("value"), "const")
            if "amount_unit" in expr:
                result *= check_unit(expr["amount_unit"])
            self.traces.append({"op": op, "value": expr["value"], "reason": expr["reason"], "amount_unit": expr.get("amount_unit"), "result": result})
        elif op == "field":
            fid = expr.get("id")
            require(fid in self.fields and self.fields[fid]["kind"] == "money", "FIELD_REFERENCE", "field 仅能引用已计算的金额字段，禁止前向引用或循环")
            precision = expr.get("precision", "display")
            require(precision in ("display", "raw"), "FIELD_PRECISION", "field precision 只能为 display 或 raw")
            for record in self.field_records.get(fid, set()):
                self.claim_record(record, record_keys)
            result = self.fields[fid]["raw_yuan"] if precision == "raw" else self.fields[fid]["value"] * UNITS[self.context["amount_unit"]]
            self.traces.append({"op": op, "id": fid, "precision": precision, "result_yuan": result})
        elif op == "table_sum":
            result = self.table_sum(expr, field, record_keys)
        elif op in ("sum", "add", "sub", "mul", "div"):
            args = expr.get("args")
            require(isinstance(args, list) and args, "EXPRESSION_ARGS", f"{op} 需要非空 args")
            require(op in ("sum", "add", "mul") or len(args) == 2, "EXPRESSION_ARGS", f"{op} 必须恰好两个参数")
            values = [self.expression(a, field, depth + 1, record_keys) for a in args]
            if op in ("sum", "add"):
                result = sum(values, Decimal(0))
            elif op == "sub":
                result = values[0] - values[1]
            elif op == "mul":
                result = Decimal(1)
                for v in values:
                    result *= v
            else:
                require(values[1] != 0, "DIVISION_ZERO", "除数为零")
                result = values[0] / values[1]
            self.traces.append({"op": op, "args": values, "result": result})
        else:
            fail("EXPRESSION_OP", f"不允许的表达式操作: {op}")
        require(result.is_finite() and abs(result) <= MAX_AMOUNT, "NUMBER_RANGE", "计算结果非有限或超限")
        return result

    def claim_record(self, record, record_keys):
        require(record not in record_keys, "RECORD_DOUBLE_COUNT", f"同一字段表达式重复纳入业务记录，不能双计: {record}")
        record_keys.add(record)

    def table_sum(self, expr, field, record_keys):
        s = self.source(expr.get("source"))
        reader, sheet = s["reader"], expr.get("sheet")
        rows, cols = reader.dimensions(sheet)
        hr = expr.get("header_row")
        require(type(hr) is int and 1 <= hr < rows, "TABLE_HEADER", "header_row 必须是有数据行的明确表头行")
        headers = expr.get("headers")
        unique = expr.get("unique_by")
        namespace = expr.get("record_namespace", "default")
        require(isinstance(namespace, str) and namespace.strip() and len(namespace) <= 128,
                "RECORD_NAMESPACE", "record_namespace 必须是明确的非空业务域名称")
        require(isinstance(headers, dict) and headers, "TABLE_HEADER", "table_sum 必须指定 headers 校验表头")
        require(isinstance(unique, list) and unique, "TABLE_KEY", "table_sum 必须指定非空 unique_by 主键列")
        used = [expr.get(k) for k in ("amount_column", "date_column", "taxpayer_column")] + unique
        require(all(isinstance(c, str) and re.fullmatch(r"[A-Z]{1,3}", c) for c in used), "TABLE_COLUMN", "table_sum 列必须是大写 Excel 列字母")
        require(len(set(unique)) == len(unique), "TABLE_KEY", "unique_by 存在重复列")
        require(all(c in headers for c in used), "TABLE_HEADER", "金额、日期、主体和主键列均需明确表头校验")
        for col, expected in headers.items():
            require(isinstance(col, str) and re.fullmatch(r"[A-Z]{1,3}", col), "TABLE_COLUMN", "headers 列格式无效")
            require(str(reader.value(sheet, f"{col}{hr}")) == str(expected), "TABLE_HEADER", f"表头不符: {sheet}!{col}{hr}")
        unit = expr.get("amount_unit", s["meta"]["context"]["amount_unit"])
        multiplier = check_unit(unit)
        start, end = self.period_for(field)
        require(s["start"] <= start and s["end"] >= end, "TABLE_PERIOD", "来源数据窗口不覆盖字段取数期间")
        total, seen, included = Decimal(0), set(), 0
        from openpyxl.utils.cell import get_column_letter
        for row in range(hr + 1, rows + 1):
            # Truly blank rows can be ignored; partially populated rows cannot.
            row_values = [reader.raw(sheet, f"{get_column_letter(c)}{row}")[0] for c in range(1, cols + 1)]
            if all(v is None or v == "" for v in row_values):
                continue
            taxid = reader.value(sheet, f"{expr['taxpayer_column']}{row}")
            require(str(taxid).strip() == self.context["taxpayer_id"], "TABLE_TAXPAYER", f"来源表含其他主体，拒绝混用: {sheet} 第 {row} 行")
            dt = iso_date(reader.value(sheet, f"{expr['date_column']}{row}"), f"{sheet} 第 {row} 行日期")
            require(s["start"] <= dt <= s["end"], "SOURCE_DATE_WINDOW", f"来源行日期超出已声明数据窗口: {sheet} 第 {row} 行")
            key = tuple(str(reader.value(sheet, f"{c}{row}")).strip() for c in unique)
            require(key not in seen, "TABLE_DUPLICATE", f"来源存在重复主键，不能静默去重: {sheet} 第 {row} 行 {key}")
            seen.add(key)
            raw = reader.value(sheet, f"{expr['amount_column']}{row}")
            value = amount(raw, f"{sheet} 第 {row} 行金额") * multiplier
            take = start <= dt <= end
            self.traces.append({"op": "table_row", "source": expr["source"], "sheet": sheet, "row": row, "key": key,
                                "record_namespace": namespace, "date": dt, "raw": serial(raw), "amount_unit": unit, "value_yuan": value, "included": take})
            if take:
                self.claim_record((namespace, key), record_keys)
                total += value
                included += 1
        require(included > 0, "TABLE_NO_ROWS", "目标期间没有可用明细，不把无数据当零申报；真实零值请提供明确零额记录")
        self.traces.append({"op": "table_sum", "source": expr["source"], "sheet": sheet, "included_rows": included, "result_yuan": total})
        return total

    def field(self, field, used):
        require(isinstance(field, dict), "FIELD_SCHEMA", "字段必须为对象")
        fid = field.get("id")
        require(isinstance(fid, str) and fid and fid not in self.fields, "FIELD_ID", "字段 id 缺少或重复")
        require(isinstance(field.get("label"), str) and field["label"].strip(), "FIELD_LABEL", f"字段缺少 label: {fid}")
        require(field.get("required") is True, "FIELD_REQUIRED", "首版仅接受 required:true 的明确字段；不静默跳过可选字段")
        scope = field.get("scope", {})
        require(isinstance(scope, dict) and all(k in self.context and self.context[k] == v for k, v in scope.items()), "FIELD_SCOPE", f"字段 scope 与当前任务不符: {fid}")
        kind = field.get("kind", "money")
        require(kind in ("money", "text"), "FIELD_KIND", "字段 kind 只能为 money 或 text")
        target = field.get("target", {})
        sheet, address = target.get("sheet"), cell_address(target.get("cell"))
        self.template.dimensions(sheet)
        key = (sheet, address)
        require(key not in used, "TARGET_DUPLICATE", f"同一目标单元格被重复写入: {sheet}!{address}")
        used.add(key)
        ws = self.template.workbook[sheet]
        c = ws[address]
        require(c.data_type != "f", "TARGET_FORMULA", f"目标是公式单元格，不能覆盖: {sheet}!{address}")
        from openpyxl.cell.cell import MergedCell
        require(not isinstance(c, MergedCell), "TARGET_MERGED", "目标为合并区域非首格")
        anchor = target.get("anchor", {})
        aa = cell_address(anchor.get("cell"))
        from openpyxl.utils.cell import coordinate_to_tuple
        tr, tc = coordinate_to_tuple(address)
        ar, ac = coordinate_to_tuple(aa)
        require((tr, tc) != (ar, ac) and abs(tr - ar) + abs(tc - ac) <= 4,
                "TARGET_ANCHOR", "目标锚点必须是距离不超过 4 格的相邻标签，不能就是目标格")
        require(isinstance(anchor.get("text"), str) and anchor["text"], "TARGET_ANCHOR", "缺少目标相邻锚点标签文本")
        require(self.template.value(sheet, aa) == anchor["text"], "TARGET_ANCHOR", f"目标标签锚点不匹配，疑似模板版本变化: {sheet}!{aa}")
        start_trace = len(self.traces)
        record_keys = set()
        expr = field.get("expression")
        if kind == "text":
            require(isinstance(expr, dict) and expr.get("op") == "ref", "TEXT_EXPRESSION", "text 字段仅能使用有来源单元格的 ref")
            s = self.source(expr.get("source"))
            raw = s["reader"].value(expr.get("sheet"), cell_address(expr.get("cell")))
            require(isinstance(raw, (str, date, datetime)), "TEXT_SOURCE_TYPE", "文本字段须来自文本或日期单元格；数字型税号可能已丢失精度或前导零")
            value = context_text(raw, self.context_bindings.get(key, {}).get("role"))
            require(not value.startswith(("=", "+", "-", "@")), "TEXT_FORMULA", "文本可能触发 Excel 公式，拒绝写入")
            require(len(value) <= 32767, "TEXT_LIMIT", "文本超过 Excel 单元格长度")
            self.traces.append({"op": "text_ref", "source": expr["source"], "sheet": expr["sheet"], "cell": expr["cell"], "value": value})
            result = {"id": fid, "label": field["label"], "kind": kind, "value": value}
        else:
            raw_yuan = self.expression(expr, field, record_keys=record_keys)
            value = (raw_yuan / UNITS[self.context["amount_unit"]]).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            require(abs(value) <= Decimal("9999999999999.99"), "EXCEL_PRECISION", "金额超过 Excel 可靠精度范围")
            require(Decimal(format(float(value), ".16g")).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP) == value,
                    "EXCEL_PRECISION", "金额转换为 Excel 数值会丢失分位精度")
            result = {"id": fid, "label": field["label"], "kind": kind, "raw_yuan": raw_yuan,
                      "value": value, "amount_unit": self.context["amount_unit"], "basis": field.get("basis", "current")}
        if key in self.context_bindings:
            binding = self.context_bindings[key]
            require(kind == "text" and value == binding["expected"], "TEMPLATE_CONTEXT_MISMATCH",
                    f"身份/期间目标必须以 text 字段写入确切任务上下文: {binding['role']}")
        if c.value is not None and c.value != "":
            same = str(c.value) == str(value)
            if kind == "money":
                try:
                    same = amount(c.value, "目标原值") == value
                except Blocked:
                    same = False
            if not same:
                require("expected_existing" in field and str(serial(c.value)) == str(field["expected_existing"]),
                        "TARGET_OCCUPIED", f"目标已有不同值，须准确声明 expected_existing 才能更新: {sheet}!{address}")
        result.update({"target": {"sheet": sheet, "cell": address}, "existing": serial(c.value), "trace": self.traces[start_trace:]})
        self.fields[fid] = result
        self.field_records[fid] = record_keys
        result["record_keys"] = [{"namespace": namespace, "key": list(key)} for namespace, key in sorted(record_keys)]
        self.audit["fields"].append(result)

    def evaluate(self):
        fields = self.job.get("fields")
        require(isinstance(fields, list) and fields, "FIELDS_MISSING", "没有待填字段")
        used = set()
        issues = self.audit["issues"]
        ids = set()
        for field in fields:
            try:
                require(isinstance(field, dict), "FIELD_SCHEMA", "字段必须为对象")
                require(field.get("id") not in ids, "FIELD_ID", "字段 id 重复")
                ids.add(field.get("id"))
                self.field(field, used)
            except Blocked as e:
                issues.append({"field": field.get("id") if isinstance(field, dict) else None, "code": e.code, "message": e.message})
        mandatory = self.job.get("mandatory_targets")
        require(isinstance(mandatory, list) and mandatory, "MANDATORY_TARGETS", "必须显式列出完整 mandatory_targets 清单")
        required_targets = set()
        for target in mandatory:
            key = (target.get("sheet"), cell_address(target.get("cell")))
            require(key not in required_targets, "MANDATORY_DUPLICATE", "mandatory_targets 重复")
            required_targets.add(key)
        successful = {(f["target"]["sheet"], f["target"]["cell"]) for f in self.fields.values()}
        for key, binding in self.context_bindings.items():
            if binding["blank"] and key not in successful:
                issues.append({"code": "TEMPLATE_CONTEXT_MISSING", "message": f"模板身份/期间空格必须由 text 来源字段完成: {binding['role']} ({key[0]}!{key[1]})"})
        for key in required_targets - successful:
            issues.append({"code": "MANDATORY_MISSING", "message": f"必填目标未完成: {key[0]}!{key[1]}"})
        require(used <= required_targets, "MANDATORY_SCOPE", "所有写入目标均必须列入 mandatory_targets，不能含清单外写入")
        checks = self.job.get("checks", [])
        require(isinstance(checks, list), "CHECK_SCHEMA", "checks 必须为列表")
        self.audit["checks"] = []
        for check in checks:
            try:
                require(isinstance(check.get("id"), str) and check["id"], "CHECK_ID", "check 缺少 id")
                left = self.expression(check.get("left"), check)
                right = self.expression(check.get("right"), check)
                tolerance = amount(check.get("tolerance", "0.00"), "容差") * UNITS[self.context["amount_unit"]]
                require(tolerance >= 0, "CHECK_TOLERANCE", "容差不能为负")
                diff = left - right
                ok = abs(diff) <= tolerance
                self.audit["checks"].append({"id": check["id"], "left_yuan": left, "right_yuan": right, "difference_yuan": diff, "tolerance_yuan": tolerance, "passed": ok})
                require(ok, "RECONCILIATION", f"勾稽不一致: {check['id']}")
            except Blocked as e:
                issues.append({"check": check.get("id"), "code": e.code, "message": e.message})
        comparisons = self.job.get("prefill_comparisons", [])
        require(isinstance(comparisons, list), "PREFILL_SCHEMA", "prefill_comparisons 必须为列表")
        for compare in comparisons:
            try:
                fid = compare.get("field")
                require(fid in self.fields and self.fields[fid]["kind"] == "money", "PREFILL_FIELD", "预填对比字段缺失或不是金额")
                basis = self.fields[fid]["basis"]
                require(compare.get("basis", basis) == basis, "PREFILL_BASIS", "预填对比 basis 必须与被比字段一致")
                actual = self.expression(compare.get("expression"), {**compare, "basis": basis})
                tol = amount(compare.get("tolerance", "0.00"), "容差") * UNITS[self.context["amount_unit"]]
                require(tol >= 0, "CHECK_TOLERANCE", "容差不能为负")
                expected = self.fields[fid]["value"] * UNITS[self.context["amount_unit"]]
                ok = abs(expected - actual) <= tol
                self.audit["checks"].append({"id": compare.get("id", fid), "type": "prefill_comparison", "field": fid, "expected_yuan": expected, "actual_yuan": actual, "tolerance_yuan": tol, "passed": ok})
                require(ok, "PREFILL_MISMATCH", f"预填值存在差异: {fid}")
            except Blocked as e:
                issues.append({"comparison": compare.get("id"), "code": e.code, "message": e.message})

    def save(self, output):
        wb = self.template.workbook
        # Recheck all hashes at the mutation boundary, including the template.
        for item in self.audit["inputs"]:
            verify_hash(Path(item["path"]), item["sha256"])
        for field in self.fields.values():
            c = wb[field["target"]["sheet"]][field["target"]["cell"]]
            c.value = field["value"]
            # Keep original template number formats, styles and other formulas.
        from openpyxl.workbook.properties import CalcProperties
        wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
        require(not output.exists(), "OUTPUT_EXISTS", "输出文件已存在，拒绝覆盖")
        wb.save(output)
        self.audit["outputs"].append({"file": output.name, "sha256": digest(output)})
        self.audit["formula_note"] = "原有非目标公式保留，已请求 Excel 打开后重算；引擎没有计算或验证公式结果。"


def review_html(audit):
    esc = lambda value: html.escape(str(serial(value)), quote=True)
    c = audit.get("context", {})
    summary = "；".join(f"{k}: {c.get(k, '')}" for k in ("taxpayer_id", "period_start", "period_end", "form_name", "form_version", "region", "amount_unit"))
    source_names = {item["id"]: Path(item["path"]).name for item in audit.get("inputs", [])}

    def trace_text(trace):
        op = trace.get("op")
        filename = source_names.get(trace.get("source"), trace.get("source", ""))
        position = f"{filename} / {trace.get('sheet', '')}"
        if op == "ref":
            result = f"取数 {position}!{trace['cell']}：原值 {trace['raw']} {trace['amount_unit']} → {serial(trace['result_yuan'])} 元"
            scope = trace.get("period_scope")
            if scope:
                anchor = scope["anchor"]
                result += f"；期间 {scope['period_start']} 至 {scope['period_end']}；已核对表头 {anchor['sheet']}!{anchor['cell']}：{anchor['observed']}"
            return result
        if op == "text_ref":
            return f"文本取数 {position}!{trace['cell']}：{trace['value']}"
        if op == "table_row":
            included = "计入" if trace["included"] else "不在目标期间，未计入"
            return (f"{position} 第 {trace['row']} 行：日期 {serial(trace['date'])}，业务键 {trace['key']}；"
                    f"原值 {trace['raw']} {trace['amount_unit']} → {serial(trace['value_yuan'])} 元；{included}")
        if op == "table_sum":
            return f"汇总 {position}：计入 {trace['included_rows']} 行，合计 {serial(trace['result_yuan'])} 元"
        if op == "const":
            unit = trace.get("amount_unit") or "（系数或元值，按表达式用途）"
            return f"常量 {trace['value']} {unit}；依据/用途：{trace['reason']}；运算值 {serial(trace['result'])}"
        if op == "field":
            precision = "未舍入原值" if trace.get("precision") == "raw" else "已填表两位值"
            return f"引用字段 {trace['id']}（{precision}）：{serial(trace['result_yuan'])} 元"
        names = {"sum": "求和", "add": "相加", "sub": "相减", "mul": "相乘", "div": "相除"}
        if op in names:
            values = ", ".join(str(serial(value)) for value in trace["args"])
            return f"{names[op]} [{values}] → {serial(trace['result'])}（金额按元参与计算）"
        return json.dumps(trace, ensure_ascii=False, default=serial)

    rows = []
    for field in audit.get("fields", []):
        unit = field.get("amount_unit", "")
        values = (field["label"], field["target"]["sheet"] + "!" + field["target"]["cell"], f"{serial(field['value'])} {unit}".rstrip())
        rows.append("<tr>" + "".join(f"<td>{esc(value)}</td>" for value in values) + "</tr>")
        trace = field.get("trace", [])
        items = "".join(f"<li>{esc(trace_text(step))}</li>" for step in trace)
        if field.get("kind") == "money":
            items += f"<li>{esc('结果：' + str(serial(field['raw_yuan'])) + ' 元 → 按 ' + unit + ' 换算并四舍五入为 ' + str(serial(field['value'])))}</li>"
        rows.append(f'<tr><td colspan="3"><details><summary>查看取数与计算（{len(trace)} 步）</summary><ol>{items}</ol></details></td></tr>')
    rows = "".join(rows)
    check_rows = []
    for check in audit.get("checks", []):
        left = check.get("left_yuan", check.get("expected_yuan"))
        right = check.get("right_yuan", check.get("actual_yuan"))
        difference = check.get("difference_yuan")
        if difference is None and left is not None and right is not None:
            difference = Decimal(str(left)) - Decimal(str(right))
        values = (check.get("id", ""), left, right, difference, check.get("tolerance_yuan", ""), "通过" if check.get("passed") else "不通过")
        check_rows.append("<tr>" + "".join(f"<td>{esc(value)}</td>" for value in values) + "</tr>")
    checks = ('<table><tr><th>检查项目</th><th>左侧/本次（元）</th><th>右侧/预填（元）</th><th>差额（元）</th><th>容差（元）</th><th>结果</th></tr>'
              + "".join(check_rows) + "</table>") if check_rows else "<p>未配置自动勾稽检查，须人工核对。</p>"
    issues = "".join(f"<li><b>{esc(i['code'])}</b> {esc(i['message'])}</li>" for i in audit.get("issues", [])) or "<li>自动检查未发现阻断项，仍须由经办会计复核。</li>"
    return """<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>纳税申报填表复核</title><style>body{font:16px/1.6 system-ui,sans-serif;max-width:1100px;margin:36px auto;padding:0 24px;color:#172033}h1{font-size:26px}table{border-collapse:collapse;width:100%}th,td{border:1px solid #ccd3df;text-align:left;padding:9px}aside{background:#fff5d6;padding:16px}li{margin:10px 0}code,td{overflow-wrap:anywhere}summary{cursor:pointer;color:#2252a0}details{font-size:14px}ol{padding-left:24px}</style><h1>纳税申报填表复核</h1>""" + f"<aside><b>{esc(audit['status'])}</b><br>这是待人工复核草稿。未登录税务系统、未提交申报、未缴款。<br>运行结果：{esc(audit['outcome'])}</aside><p>{esc(summary)}</p><h2>检查结果</h2><ul>{issues}</ul><h2>字段与来源</h2><table><tr><th>字段</th><th>目标</th><th>值</th></tr>{rows}</table><h2>勾稽与预填对比</h2>{checks}<p>完整取数、运算、行级记录及输入输出 SHA256 同时保存在 audit.json。原有公式由 Excel 打开后重算，引擎未计算公式。</p></html>"


def run(job_path, out):
    require(not out.exists() and not out.is_symlink(), "OUTPUT_EXISTS", "--out 必须是尚不存在的新目录，不允许覆盖")
    jp = input_path(Path.cwd(), str(job_path))
    out = out.expanduser().resolve()
    require(out != jp and out not in jp.parents, "OUTPUT_INPUT", "输出目录不能包含或覆盖输入任务")
    audit = {"schema_version": 1, "engine_version": VERSION, "status": STATUS, "outcome": "BLOCKED",
             "created_at": datetime.now(timezone.utc).isoformat(), "context": {}, "inputs": [], "fields": [], "checks": [], "issues": [], "outputs": [],
             "human_review_required": True, "filing_performed": False}
    engine, created = None, False
    try:
        require(jp.suffix.lower() == ".json", "JOB_FORMAT", "任务必须为 JSON 文件")
        with jp.open(encoding="utf-8-sig") as f:
            job = json.load(f, parse_float=Decimal)
        require(isinstance(job, dict), "JOB_FORMAT", "任务顶层必须为 JSON 对象")
        audit["context"] = job.get("context", {})
        audit["inputs"].append({"id": "job", "path": str(jp), "sha256": digest(jp)})
        # Resolve every declared input before creating any output directory.
        declared = [s.get("path") for s in job.get("sources", []) if isinstance(s, dict)]
        declared.append(job.get("template", {}).get("path"))
        for value in declared:
            if isinstance(value, str):
                p = input_path(jp.parent, value)
                require(out != p and out not in p.parents, "OUTPUT_INPUT", "输出目录不能包含任何输入文件")
        with localcontext() as ctx:
            ctx.prec = 50
            engine = Engine(job, jp.parent, audit)
            engine.validate()
            engine.evaluate()
            if not audit["issues"]:
                out.mkdir(parents=True, exist_ok=False)
                created = True
                engine.save(out / "filled.xlsx")
                audit["outcome"] = "GENERATED_FOR_REVIEW"
    except Blocked as e:
        audit["issues"].append({"code": e.code, "message": e.message})
    except (json.JSONDecodeError, UnicodeDecodeError) as e:
        audit["issues"].append({"code": "JOB_PARSE", "message": str(e)})
    except Exception as e:
        audit["issues"].append({"code": "ENGINE_ERROR", "message": f"{type(e).__name__}: {e}"})
    finally:
        if engine:
            audit["calculation_trace"] = engine.traces
            for s in engine.sources.values():
                if s["reader"].workbook:
                    s["reader"].workbook.close()
            if hasattr(engine, "template"):
                engine.template.workbook.close()
    if audit["issues"]:
        audit["outcome"] = "BLOCKED"
        partial = out / "filled.xlsx"
        if created and partial.exists():
            partial.unlink()
        audit["outputs"] = []
    if not created:
        out.mkdir(parents=True, exist_ok=False)
        created = True
    review = out / "review.html"
    with review.open("x", encoding="utf-8") as f:
        f.write(review_html(audit))
    audit["outputs"].append({"file": review.name, "sha256": digest(review)})
    dump_json(out / "audit.json", audit)
    print(json.dumps({"status": STATUS, "outcome": audit["outcome"], "issues": len(audit["issues"]), "output_directory": str(out),
                      "audit_sha256": digest(out / "audit.json")}, ensure_ascii=False))
    return 2 if audit["issues"] else 0


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--version", action="version", version=VERSION)
    sub = parser.add_subparsers(dest="command", required=True)
    sp = sub.add_parser("scan", help="安全扫描 CSV/XLSX，不读取公式计算值")
    sp.add_argument("directory", type=Path)
    sp.add_argument("--output", required=True, type=Path)
    rp = sub.add_parser("run", help="填充模板副本并生成复核报告，始终为草稿")
    rp.add_argument("job", type=Path)
    rp.add_argument("--out", required=True, type=Path)
    args = parser.parse_args(argv)
    try:
        import openpyxl  # noqa: F401
        if args.command == "scan":
            return scan(args.directory, args.output.expanduser().resolve())
        return run(args.job, args.out)
    except ModuleNotFoundError as e:
        print(json.dumps({"status": "BLOCKED", "code": "DEPENDENCY_MISSING", "message": f"缺少依赖 {e.name}；仅需 openpyxl，请在当前隔离 Python 环境安装"}, ensure_ascii=False), file=sys.stderr)
        return 2
    except Blocked as e:
        print(json.dumps({"status": "BLOCKED", "code": e.code, "message": e.message}, ensure_ascii=False), file=sys.stderr)
        return 2
    except Exception as e:
        print(json.dumps({"status": "BLOCKED", "code": "ENGINE_ERROR", "message": f"{type(e).__name__}: {e}"}, ensure_ascii=False), file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
