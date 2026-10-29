#!/usr/bin/env python3
"""Profit-statement -> EIT rules -> source-traced official-line working draft.

The Agent writes case.json from the user's files and explicit confirmation.
This module does not establish tax facts from silence, submit returns or claim
that the reconstructed template is a tax-bureau Excel import format.
"""
from __future__ import annotations

import argparse
from contextlib import redirect_stdout
from copy import deepcopy
import csv
from datetime import date, datetime
from decimal import Decimal
import html
import io
import json
from pathlib import Path
import re
import sys

import eit_rules
from fact_sources import resolve_sources
import statement_adapter
import tax_fill

FORM_DIR = Path(__file__).resolve().parents[1] / "assets"
STATUS = "DRAFT_REVIEW_REQUIRED"
UNITS = {"元": Decimal(1), "千元": Decimal(1000), "万元": Decimal(10000)}
FINANCIAL_KEYS = [f[0] for f in eit_rules.FINANCIAL_FIELDS]


def need(condition, code, message):
    if not condition:
        raise tax_fill.Blocked(code, message)


def problem(code, message, **extra):
    return {"code": code, "message": message, **extra}


def read_json(path):
    with path.open(encoding="utf-8-sig") as stream:
        result = json.load(stream, parse_float=Decimal)
    need(isinstance(result, dict), "JSON_SCHEMA", "JSON 顶层必须为对象")
    return result


def text_evidence(value):
    return isinstance(value, str) and bool(value.strip())


def line_order(number):
    return tuple(int(part) for part in number.split("."))


def literal_csv(value):
    """CSV has no cell type; quote formula-like descriptive text as literal."""
    value = str(value)
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def source_context(case, case_path):
    context = case.get("context", {})
    need(isinstance(context, dict), "CONTEXT_SCHEMA", "context 必须为对象")
    for key in ("taxpayer_id", "taxpayer_name", "period_start", "period_end", "period_kind", "form_name", "form_version", "region", "currency", "amount_unit"):
        need(text_evidence(context.get(key)), "CONTEXT_MISSING", f"缺少明确上下文 {key}")
    need(context["currency"] == "CNY" and context["amount_unit"] == "元", "CONTEXT_CURRENCY_UNIT", "本业务链路仅输出人民币元，不自动换汇")
    need(not any(context[key].lstrip().startswith(("=", "+", "-", "@")) for key in ("taxpayer_id", "taxpayer_name")),
         "IDENTITY_LITERAL_UNSAFE", "主体税号或名称不能作为可能执行的表格公式写入")
    need(case.get("return_id") in ("A200000", "A100000"), "RETURN_UNSUPPORTED", "本链路只适配 A200000 和 A100000")
    need(text_evidence(case.get("confirmation_evidence")), "CONFIRMATION_REQUIRED", "需要本次主体、所属期、材料范围和口径的真实确认记录")
    start = tax_fill.iso_date(context["period_start"], "context.period_start")
    end = tax_fill.iso_date(context["period_end"], "context.period_end")
    need(start.year == end.year and start <= end <= date.today(), "CONTEXT_PERIOD", "只处理同年已结束的明确所属期")
    statement = case.get("income_statement", {})
    need(isinstance(statement, dict), "STATEMENT_SCHEMA", "income_statement 必须为对象")
    for key in ("path", "sheet", "accounting_standard", "presentation", "source_period_start", "source_period_end", "period_evidence"):
        need(text_evidence(statement.get(key)), "STATEMENT_CONTEXT", f"利润表缺少明确 {key}")
    need(statement["accounting_standard"] in ("asbe", "small_asbe"), "ACCOUNTING_STANDARD", "需要明确 asbe 或 small_asbe")
    need(statement["presentation"] in ("software_export", "tax_filing"), "PRESENTATION", "需要明确软件导出或税务财报语境")
    ss = tax_fill.iso_date(statement["source_period_start"], "source_period_start")
    se = tax_fill.iso_date(statement["source_period_end"], "source_period_end")
    need(ss == date(end.year, 1, 1) and se == end, "STATEMENT_PERIOD", "源利润表累计窗口必须从所属年度1月1日至本期末，不能拿本季或其他期间替代")
    if "amount_unit" in statement:
        need(statement["amount_unit"] in UNITS, "SOURCE_UNIT", "来源单位必须为元、千元或万元")
        need(text_evidence(statement.get("unit_evidence")), "UNIT_EVIDENCE_REQUIRED", "显式来源单位需要实际表头或用户确认依据")
    input_dir = Path(case.get("input_dir", "")).expanduser()
    need(text_evidence(case.get("input_dir")), "INPUT_DIRECTORY", "需要明确本次资料目录")
    if not input_dir.is_absolute():
        input_dir = case_path.parent / input_dir
    input_dir = input_dir.resolve()
    need(input_dir.is_dir(), "INPUT_DIRECTORY", "资料目录不存在")
    source = tax_fill.input_path(input_dir, statement["path"])
    need(source.is_relative_to(input_dir), "STATEMENT_OUTSIDE_INPUT", "利润表必须位于本次指定资料目录内")
    return context, statement, input_dir, source


def metadata_checks(path, statement, context, selected_sheet):
    """Check explicit visible conflicts, without guessing ambiguous dates."""
    reader = tax_fill.Source(path)
    evidence, anchors = [], []
    try:
        sheet = statement["sheet"]
        reader.dimensions(sheet)
        rows, cols = reader.dimensions(sheet)
        mapped_rows = [int(re.search(r"\d+$", m["source"]["label_cell"]).group()) for m in selected_sheet["mappings"]]
        limit = min(mapped_rows, default=min(rows + 1, 21))
        entries = []
        from openpyxl.utils.cell import get_column_letter
        for row in range(1, min(limit, 21)):
            for col in range(1, cols + 1):
                cell = f"{get_column_letter(col)}{row}"
                value, kind = reader.raw(sheet, cell)
                if value is not None and kind not in ("f", "e"):
                    entries.append((row, col, cell, value))
        expected_start, expected_end = statement["source_period_start"], statement["source_period_end"]
        end_date = date.fromisoformat(expected_end)
        for row, col, cell, raw in entries:
            value = raw.date().isoformat() if isinstance(raw, datetime) else raw.isoformat() if isinstance(raw, date) else str(tax_fill.serial(raw)).strip()
            normalized = statement_adapter.norm(value)
            currency = re.search(r"(?:币种|货币|Currency)[:：]?(USD|EUR|HKD|美元|欧元|港币|港元|JPY|日元|GBP|英镑)", normalized, re.I)
            foreign_unit = re.search(r"(?:金额单位|单位)[:：]?(?:千|万)?(美元|欧元|港币|港元|日元|英镑)", normalized)
            need(not currency and not foreign_unit, "FOREIGN_CURRENCY_METADATA", f"利润表 {sheet}!{cell} 明确为外币，当前链路不换汇")
            name_match = re.fullmatch(r"(?:编制单位|纳税人名称|企业名称|公司名称)[:：](.+)", value)
            if name_match:
                stated = name_match.group(1).strip()
                need(statement_adapter.norm(stated) == statement_adapter.norm(context["taxpayer_name"]),
                     "TAXPAYER_NAME_MISMATCH", f"利润表 {sheet}!{cell} 企业名称与本次主体冲突")
                evidence.append({"type": "visible_name", "cell": cell, "raw": value, "matched_context_key": "taxpayer_name"})
                anchors.append({"sheet": sheet, "cell": cell, "expected": value})
            # A label and its value may occupy adjacent cells.
            if normalized.rstrip(":：") in ("编制单位", "纳税人名称", "企业名称", "公司名称", "纳税人识别号", "统一社会信用代码"):
                neighbor_cell = f"{get_column_letter(col + 1)}{row}"
                neighbor, kind = reader.raw(sheet, neighbor_cell)
                if neighbor is not None and kind not in ("f", "e"):
                    key = "taxpayer_id" if normalized.rstrip(":：") in ("纳税人识别号", "统一社会信用代码") else "taxpayer_name"
                    need(isinstance(neighbor, str), "IDENTITY_TEXT_REQUIRED", f"身份元数据 {neighbor_cell} 必须为文本，数字税号可能已经失真")
                    need(statement_adapter.norm(neighbor) == statement_adapter.norm(context[key]), "IDENTITY_METADATA_MISMATCH", f"利润表 {neighbor_cell} 与 {key} 冲突")
                    evidence.append({"type": "visible_identity", "cell": neighbor_cell, "raw": neighbor, "matched_context_key": key})
                    anchors.append({"sheet": sheet, "cell": neighbor_cell, "expected": neighbor})
            # Ignore explicitly comparative date headers; they are valid past data.
            if any(token in normalized for token in ("上年", "上期", "同期", "比较")):
                continue
            stripped = re.sub(r"^(?:日期|报表日期|会计期间|所属期|会计年度|报表期间)[:：]?", "", normalized)
            dates = re.findall(r"20\d{2}-\d{2}-\d{2}", stripped)
            if len(dates) == 2 and (stripped == "至".join(dates) or stripped == "~".join(dates)):
                need(dates == [expected_start, expected_end], "PERIOD_METADATA_CONFLICT", f"利润表 {cell} 明确起止期与累计窗口冲突")
                evidence.append({"type": "visible_period_range", "cell": cell, "raw": value})
            elif re.fullmatch(r"20\d{2}-\d{2}-\d{2}", stripped):
                need(stripped in (expected_start, expected_end), "PERIOD_METADATA_CONFLICT", f"利润表 {cell} 日期与累计窗口端点冲突")
                evidence.append({"type": "visible_period_endpoint", "cell": cell, "raw": value})
            else:
                ym = re.fullmatch(r"(20\d{2})(?:年|-|/)(\d{1,2})月?", stripped)
                if ym:
                    need((int(ym[1]), int(ym[2])) == (end_date.year, end_date.month), "PERIOD_METADATA_CONFLICT", f"利润表 {cell} 年月与本期末冲突")
                    evidence.append({"type": "visible_year_month", "cell": cell, "raw": value, "scope": "end_month_only"})
        declared = statement.get("identity_anchors", [])
        need(isinstance(declared, list), "IDENTITY_ANCHOR_SCHEMA", "identity_anchors 必须为列表")
        for anchor in declared:
            key = anchor.get("expected_context_key")
            need(key in ("taxpayer_id", "taxpayer_name", "period_end"), "IDENTITY_ANCHOR_SCHEMA", "身份锚点只支持主体税号、名称和期末日期")
            anchor_sheet = anchor.get("sheet", sheet)
            cell = tax_fill.cell_address(anchor.get("cell"))
            value = reader.value(anchor_sheet, cell)
            if key == "period_end":
                actual = tax_fill.iso_date(value, "period_end anchor").isoformat()
                same = actual == context[key]
            else:
                need(isinstance(value, str), "IDENTITY_TEXT_REQUIRED", "税号与企业名称锚点须为文本")
                same = statement_adapter.norm(value) == statement_adapter.norm(context[key])
            need(same, "IDENTITY_ANCHOR_MISMATCH", f"利润表 {anchor_sheet}!{cell} 与 {key} 不一致")
            evidence.append({"type": "explicit_anchor", "sheet": anchor_sheet, "cell": cell, "raw": tax_fill.serial(value), "matched_context_key": key})
            anchors.append({"sheet": anchor_sheet, "cell": cell, "expected": str(tax_fill.serial(value))})
        # Absence of anchors is not silently promoted to verified identity.
        return {"identity_evidence": "visible_anchors_and_human_declaration" if anchors else "human_declaration_only",
                "verified_context_keys": sorted({e["matched_context_key"] for e in evidence if "matched_context_key" in e}),
                "period_window_evidence": "human_declaration_with_visible_conflict_checks",
                "declaration": statement["period_evidence"], "observations": evidence, "anchors": anchors}
    finally:
        if reader.workbook:
            reader.workbook.close()


def select_financials(scan, source, statement, case):
    file_result = next((f for f in scan["files"] if Path(f["path"]) == source), None)
    need(file_result is not None, "STATEMENT_NOT_SCANNED", "指定利润表未被扫描；检查格式、重复副本或路径")
    need(not file_result["issues"], "STATEMENT_READ_ERROR", "指定利润表读取存在错误，不能绕过")
    need(not scan.get("inputs_changed"), "HASH_CHANGED", "扫描期间原始资料发生变化")
    selected = next((s for s in file_result["sheets"] if s["name"] == statement["sheet"]), None)
    need(selected is not None and selected["classification"] == "profit_statement", "STATEMENT_SHEET", "指定工作表未识别为利润表")
    structural = {"ACCOUNTING_STANDARD_CONFLICT", "HEADER_CONFLICT", "UNIT_CONFLICT", "DUPLICATE_SEMANTIC_ROW", "MISSING_PERIOD_HEADER"}
    errors = [p for p in selected["issues"] if p["code"] in structural]
    need(not errors, "STATEMENT_STRUCTURE", "利润表存在准则、单位、期间列或主行冲突：" + "; ".join(p["code"] for p in errors))
    for problem_ in scan["issues"]:
        if problem_["code"] == "DUPLICATE_FILE_CONTENT" and str(source) in (problem_.get("path"), problem_.get("first_path")):
            raise tax_fill.Blocked("DUPLICATE_FILE_CONTENT", "本次资料含相同利润表副本，请明确保留哪个原始来源")
    financials, mapped = {}, {}
    for key in FINANCIAL_KEYS:
        all_candidates = [m for m in selected["mappings"] if m["field"] == key]
        if not all_candidates:
            continue  # Only eit_rules may resolve absent rows via explicit zero facts.
        candidates = [m for m in all_candidates if m["basis"] == "ytd"]
        need(len(candidates) == 1, "YTD_MAPPING_REQUIRED", f"{key} 需要唯一、明确的本年累计列，不能以本月/本季代替或把歧义当零")
        mapping = deepcopy(candidates[0])
        effective_unit = mapping.get("amount_unit")
        supplied_unit = statement.get("amount_unit")
        if supplied_unit and effective_unit:
            need(supplied_unit == effective_unit, "UNIT_OVERRIDE_CONFLICT", "表内单位与显式单位不同，不能覆盖真实表头")
        resolved = []
        if effective_unit is None and supplied_unit:
            effective_unit = supplied_unit
            mapping["unit_confirmation"] = statement["unit_evidence"]
            resolved = [p for p in mapping["issues"] if p["code"] == "UNIT_UNCONFIRMED"]
        remaining = [p for p in mapping["issues"] if p not in resolved]
        need(not remaining and effective_unit in UNITS and mapping.get("raw_numeric") is not None,
             "UNUSABLE_YTD_SOURCE", f"{key} 源值不可用：" + "; ".join(p["code"] for p in remaining))
        numeric = tax_fill.amount(mapping["raw_numeric"], key)
        value = numeric * UNITS[effective_unit]
        mapping.update(effective_amount_unit=effective_unit, unit_multiplier=str(UNITS[effective_unit]),
                       value_yuan=format(value, "f"), resolved_unit_issues=resolved,
                       period_scope={"period_start": statement["source_period_start"], "period_end": statement["source_period_end"],
                                     "evidence": statement["period_evidence"]})
        financials[key] = {"value": format(value, "f"), "source": mapping}
        mapped[key] = mapping
    need(mapped, "NO_YTD_FINANCIALS", "未找到可取数的累计财务字段")
    identity = metadata_checks(source, statement, case["context"], selected)
    return financials, mapped, identity


def check_hashes(items):
    failures = []
    for item in items:
        path = Path(item["path"])
        if not path.is_file() or tax_fill.digest(path) != item["sha256"]:
            failures.append(problem("HASH_CHANGED", "原始输入在执行期间变化", path=str(path)))
    return failures


def build_lineage(calculation, mapped, case, identity, source, source_sha):
    lines = calculation.get("lines", {})
    def field_dependencies(number, visited=None):
        visited = set() if visited is None else set(visited)
        if number in visited:
            return set()
        visited.add(number)
        result = set()
        for item in lines[number].get("source_fields", []):
            if isinstance(item, str) and item in lines:
                result.update(field_dependencies(item, visited))
            elif isinstance(item, dict) and item.get("field") in FINANCIAL_KEYS:
                result.add(item["field"])
        if number.isdigit() and 1 <= int(number) <= len(FINANCIAL_KEYS):
            result.add(FINANCIAL_KEYS[int(number) - 1])
        if re.fullmatch(r"9\.[1-9][0-9]*", number):
            result.add("investment_income")
        return result
    return {"schema_version": 1, "return_id": case.get("return_id"), "context": case.get("context"),
            "statement_source": {"path": str(source), "sha256": source_sha, **identity},
            "financial_sources": mapped, "confirmation_evidence": case.get("confirmation_evidence"),
            "confirmed_facts": case.get("facts", {}), "filing_performed": False,
            "lines": {number: {**deepcopy(line), "financial_sources": {key: mapped[key] for key in sorted(field_dependencies(number)) if key in mapped},
                                "confirmed_facts_reference": "#/confirmed_facts"} for number, line in lines.items()}}


def write_csv(path, rows):
    with path.open("x", encoding="utf-8-sig", newline="") as stream:
        csv.writer(stream).writerows(rows)


def resolve_template(case, out, calculation, meta, template):
    """Add only verified A200000 investment-item views to a private copy."""
    extra = set(calculation["lines"]) - set(meta.get("line_targets", {})) - set(meta.get("fixed_values", {}))
    if not extra:
        return meta, template, None
    need(case["return_id"] == "A200000" and all(re.fullmatch(r"9\.[1-9][0-9]*", n) for n in extra),
         "LINE_TARGET_MISSING", "存在未适配的规则行号，不能自动生成任意未知税表目标")
    tax_fill.verify_hash(template, meta.get("template_sha256"))
    reader = tax_fill.Source(template, template=True)
    resolved = out / "resolved_template.xlsx"
    try:
        wb = reader.workbook
        need("投资收益明细" not in wb.sheetnames, "DYNAMIC_SHEET_CONFLICT", "基础模板已经有同名事项页，不能覆盖")
        ws = wb.create_sheet("投资收益明细")
        ws.append(["事项行次", "校验标签", "金额（元）", "已确认事项名称"])
        ws["F1"] = "工作底稿事项视图；非税局正式附表布局"
        updated = deepcopy(meta)
        items = case.get("facts", {}).get("investment_income_items", [])
        for target_row, number in enumerate(sorted(extra, key=line_order), 2):
            index = int(number.split(".")[1])
            need(isinstance(items, list) and index <= len(items) and isinstance(items[index-1], dict),
                 "INVESTMENT_DETAIL_REQUIRED", f"规则第 {number} 行没有对应的已确认事项")
            item, line = items[index-1], calculation["lines"][number]
            need(text_evidence(item.get("name")) and line.get("label") == item["name"] and line.get("unit") == "元",
                 "DYNAMIC_ITEM_CONFLICT", f"第 {number} 行事项名称或金额单位与已核验规则不一致")
            original = tax_fill.amount(item.get("value"), "投资收益事项").quantize(Decimal("0.01"), rounding=tax_fill.ROUND_HALF_UP)
            need(original == tax_fill.amount(line["value"], "规则事项值"), "DYNAMIC_ITEM_CONFLICT", f"第 {number} 行规则金额与事实来源不符")
            anchor = f"投资收益事项{index}"
            for address, value in ((f"A{target_row}", number), (f"B{target_row}", anchor), (f"D{target_row}", item["name"])):
                ws[address] = value
                ws[address].data_type = "s"  # Preserve formula-like names literally.
            ws[f"C{target_row}"].number_format = '#,##0.00;[Red]-#,##0.00'
            updated["line_targets"][number] = {"sheet": ws.title, "cell": f"C{target_row}",
                                                "anchor": {"cell": f"B{target_row}", "text": anchor}}
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 24
        ws.column_dimensions["C"].width = 20
        ws.column_dimensions["D"].width = 50
        wb.save(resolved)
        updated["template_sha256"] = tax_fill.digest(resolved)
        transformation = {"operation": "append_literal_investment_item_view", "base_template": str(template),
                          "base_sha256": meta["template_sha256"], "resolved_template": str(resolved),
                          "resolved_sha256": updated["template_sha256"], "added_sheet": "投资收益明细",
                          "added_lines": sorted(extra, key=line_order), "main_sheet_layout_changed": False,
                          "meaning": "工作底稿事项视图，不是税局正式附表布局"}
        tax_fill.dump_json(out / "resolved_template.json", updated)
        return updated, resolved, transformation
    finally:
        reader.workbook.close()


def make_job(case, out, calculation, meta, template, source, source_sha, mapped, identity):
    context = case["context"]
    need(meta.get("form_id") == case["return_id"], "FORM_ID", "模板 form_id 与任务不一致")
    need(meta.get("form_version") == context["form_version"] == calculation.get("form_version"), "FORM_VERSION", "规则、任务与模板版本不一致")
    tax_fill.verify_hash(template, meta.get("template_sha256"))
    lines = calculation["lines"]
    targets, fixed = meta.get("line_targets", {}), meta.get("fixed_values", {})
    need(isinstance(targets, dict) and isinstance(fixed, dict), "FORM_META", "模板必须声明 line_targets 和 fixed_values")
    reader = tax_fill.Source(template, template=True)
    try:
        for number, definition in fixed.items():
            need(number in lines, "FIXED_LINE_UNKNOWN", f"固定行 {number} 不在规则结果中")
            actual = reader.value(definition["sheet"], definition["cell"])
            need(lines[number].get("unit") == "rate" and "%" in reader.workbook[definition["sheet"]][definition["cell"]].number_format,
                 "FIXED_UNIT_CONFLICT", f"固定第 {number} 行应保持比率单位及百分比格式")
            need(tax_fill.amount(actual, "固定模板值") == tax_fill.amount(definition["value"], "固定定义") == tax_fill.amount(lines[number]["value"], "规则固定值"),
                 "FIXED_VALUE_CONFLICT", f"固定行 {number} 与模板/规则不同")
    finally:
        reader.workbook.close()
    for number in lines:
        need(number in targets or number in fixed, "LINE_TARGET_MISSING", f"实际税表第 {number} 行没有目标映射，不能称整张主表完成")
    value_csv, context_csv = out / "return_values.csv", out / "context_values.csv"
    value_rows = [["line", "label", "value", "unit", "calculation", "policy_url"]]
    row_for = {}
    for number in sorted(lines, key=line_order):
        row_for[number] = len(value_rows) + 1
        line = lines[number]
        value_rows.append([number, literal_csv(line["label"]), line["value"], line["unit"], literal_csv(line["calculation"]), literal_csv(line["rule_url"])])
    write_csv(value_csv, value_rows)
    identity_defs = meta.get("identity_fields")
    need(isinstance(identity_defs, list) and identity_defs, "IDENTITY_TARGETS", "模板没有身份/期间字段定义")
    context_rows = [["context_key", "value", "confirmation"]]
    context_row_for = {}
    for definition in identity_defs:
        key = definition["context_key"]
        need(key in context and key not in context_row_for, "IDENTITY_TARGETS", "身份字段未知或重复")
        context_row_for[key] = len(context_rows) + 1
        context_rows.append([key, context[key], case["confirmation_evidence"]])
    write_csv(context_csv, context_rows)
    source_window = {"taxpayer_id": context["taxpayer_id"], "period_start": f"{context['period_end'][:4]}-01-01",
                     "period_end": context["period_end"], "currency": "CNY", "amount_unit": "元"}
    fields = []
    for definition in identity_defs:
        key = definition["context_key"]
        fields.append({"id": f"identity_{key}", "label": definition["label"], "kind": "text", "required": True,
                       "target": definition["target"], "expression": {"op": "ref", "source": "context_values", "sheet": "CSV", "cell": f"B{context_row_for[key]}"}})
    for number in sorted(lines, key=line_order):
        if number in fixed:
            continue
        fields.append({"id": f"line_{number}", "label": lines[number]["label"], "kind": "money", "required": True, "basis": "ytd",
                       "target": targets[number], "expression": {"op": "ref", "source": "return_values", "sheet": "CSV", "cell": f"C{row_for[number]}"}})
    original_unit = next(iter(mapped.values()))["effective_amount_unit"]
    sources = [{"id": name, "path": str(path), "sha256": tax_fill.digest(path), "context": source_window, "metadata_confirmed": True,
                "identity_anchors": []} for name, path in (("return_values", value_csv), ("context_values", context_csv))]
    sources.append({"id": "original_statement", "path": str(source), "sha256": source_sha,
                    "context": {**source_window, "amount_unit": original_unit}, "metadata_confirmed": True,
                    "identity_anchors": identity["anchors"]})
    return {"schema_version": 1, "context": context, "confirmations": {"context_confirmed": True, "materials_complete": True},
            "confirmation_evidence": case["confirmation_evidence"], "sources": sources,
            "template": {"path": str(template), "sha256": meta["template_sha256"], "context_bindings": meta["context_bindings"]},
            "fields": fields, "mandatory_targets": [f["target"] for f in fields], "checks": [],
            "lineage_file": "lineage.json", "rules_calculation_file": "calculation.json"}


def html_report(result):
    esc = lambda value: html.escape(str(value), quote=True)
    issues = "".join(f"<li><b>{esc(p['code'])}</b> {esc(p.get('message',''))}</li>" for p in result["issues"]) or "<li>当前分支检查通过，仍须会计复核。</li>"
    links = "".join(f'<li><a href="{esc(path)}">{esc(name)}</a></li>' for name, path in result["outputs"].items())
    return '<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>企业所得税填写稿</title><style>body{font:16px/1.7 system-ui;max-width:1000px;margin:40px auto;padding:0 24px}aside{background:#fff2cf;padding:16px}</style>' + f'<h1>企业所得税填表结果</h1><aside>{esc(result["outcome"])}<br>真实官方行次重建的工作底稿，待会计复核。未提交申报、未缴款；未证明税局可导入。</aside><h2>需处理事项</h2><ul>{issues}</ul><h2>文件</h2><ul>{links}</ul><p>累计窗口与企业资料包含人工声明；可见元数据和锚点用于冲突检查，不能替代经济实质及完整性复核。</p></html>'


def run_case(case_path, out, forms_dir=None):
    case_path = tax_fill.input_path(Path.cwd(), str(case_path))
    out = Path(out).expanduser().absolute()
    need(not out.exists() and not out.is_symlink(), "OUTPUT_EXISTS", "输出必须为不存在的新目录")
    need(out.resolve() not in case_path.parents, "OUTPUT_INPUT", "输出目录不能包含输入任务")
    case = read_json(case_path)
    # Do not create a directory over any source material, even for blocked cases.
    declared_root = case.get("input_dir")
    if text_evidence(declared_root):
        declared = Path(declared_root).expanduser()
        declared = (case_path.parent / declared).resolve() if not declared.is_absolute() else declared.resolve()
        need(out.resolve() != declared and out.resolve() not in declared.parents, "OUTPUT_INPUT", "输出不能包含资料目录")
    out.mkdir(parents=True, exist_ok=False)
    result = {"status": STATUS, "outcome": "BLOCKED", "return_id": case.get("return_id"), "issues": [], "outputs": {},
              "human_review_required": True, "filing_performed": False, "input_hashes_unchanged": None}
    calculation = {"status": STATUS, "outcome": "BLOCKED", "lines": {}, "issues": [], "calculation_performed": False}
    lineage = {"schema_version": 1, "lines": {}, "filing_performed": False}
    hash_items = [{"path": str(case_path), "sha256": tax_fill.digest(case_path)}]
    try:
        context, statement, input_dir, source = source_context(case, case_path)
        scan = statement_adapter.scan_directory(input_dir, statement["accounting_standard"], statement["presentation"])
        tax_fill.dump_json(out / "scan.json", scan)
        result["outputs"]["原始资料识别与歧义"] = "scan.json"
        hash_items.extend({"path": f["path"], "sha256": f["sha256"]} for f in scan["files"])
        financials, mapped, identity = select_financials(scan, source, statement, case)
        need(not check_hashes(hash_items), "HASH_CHANGED", "取数后源文件变化，需重新扫描")
        facts = case.get("facts", {})
        need(isinstance(facts, dict), "FACTS_SCHEMA", "facts 必须为规则模块的已确认事实对象")
        resolved_facts, fact_manifests, fact_trace = resolve_sources(facts, input_dir, context)
        # Keep the original case file intact; lineage records the actual resolved
        # cell values, while explicit source-less facts remain declarations.
        case = {**case, "facts": resolved_facts}
        for manifest in fact_manifests:
            existing = next((item for item in hash_items if item["path"] == manifest["path"]), None)
            if existing is not None:
                need(existing["sha256"] == manifest["sha256"], "HASH_CHANGED", "专项来源在财报扫描和事实读取之间变化")
                existing.update(manifest)
            else:
                hash_items.append(manifest)
        need(not check_hashes(hash_items), "HASH_CHANGED", "专项事实读取后输入变化")
        calculation = eit_rules.calculate(case["return_id"], financials, deepcopy(resolved_facts), deepcopy(context))
        source_sha = tax_fill.digest(source)
        lineage = build_lineage(calculation, mapped, case, identity, source, source_sha)
        lineage["fact_source_trace"] = fact_trace
        lineage["fact_source_manifests"] = fact_manifests
        result["issues"].extend(calculation.get("issues", []))
        if calculation["outcome"] != "BLOCKED":
            folder = Path(forms_dir) if forms_dir is not None else FORM_DIR
            meta_path = folder / f"{case['return_id']}.json"
            template = folder / f"{case['return_id']}.xlsx"
            need(meta_path.is_file() and template.is_file(), "FORM_ASSET_MISSING", "实际税表行次模板资源尚未就绪，不能以示例模板冒充")
            meta = read_json(meta_path)
            hash_items.extend([{"path": str(meta_path), "sha256": tax_fill.digest(meta_path)}, {"path": str(template), "sha256": tax_fill.digest(template)}])
            meta, template, transformation = resolve_template(case, out, calculation, meta, template)
            if transformation:
                lineage["template_transformation"] = transformation
                hash_items.append({"path": str(template), "sha256": meta["template_sha256"]})
                result["outputs"].update({"含投资收益事项视图的私有模板副本": "resolved_template.xlsx", "副本模板行次映射": "resolved_template.json"})
            job = make_job(case, out, calculation, meta, template, source, source_sha, mapped, identity)
            tax_fill.dump_json(out / "job.json", job)
            result["outputs"].update({"派生税表数值来源": "return_values.csv", "身份与期间确认来源": "context_values.csv", "填写任务与全部目标": "job.json"})
            need(not check_hashes(hash_items), "HASH_CHANGED", "执行填写前输入变化")
            with redirect_stdout(io.StringIO()):
                code = tax_fill.run(out / "job.json", out / "fill-result")
            audit = read_json(out / "fill-result" / "audit.json")
            result["outputs"].update({"填写引擎审计": "fill-result/audit.json", "逐字段复核报告": "fill-result/review.html"})
            if code == 0 and audit["outcome"] == "GENERATED_FOR_REVIEW":
                result["outcome"] = "GENERATED_FOR_REVIEW"
                result["outputs"]["填写后的企业所得税主表复核稿"] = "fill-result/filled.xlsx"
            else:
                result["issues"].extend(audit.get("issues", []))
    except tax_fill.Blocked as error:
        result["issues"].append(problem(error.code, error.message))
    except Exception as error:
        result["issues"].append(problem("WORKFLOW_ERROR", f"{type(error).__name__}: {error}"))
    finally:
        changed = check_hashes(hash_items)
        result["input_hashes_unchanged"] = not changed
        result["input_hash_checks"] = hash_items
        result["issues"].extend(changed)
        if result["issues"]:
            result["outcome"] = "BLOCKED"
            filled = out / "fill-result" / "filled.xlsx"
            if filled.is_file():
                filled.unlink()
                result["issues"].append(problem("DRAFT_WITHDRAWN", "源资料或填写状态变化，填写稿已撤下；子审计保留当时运行记录，不代表有效交付"))
            result["outputs"].pop("填写后的企业所得税主表复核稿", None)
        if calculation.get("calculation_performed") is False:
            calculation["issues"] = result["issues"]
        tax_fill.dump_json(out / "calculation.json", calculation)
        tax_fill.dump_json(out / "lineage.json", lineage)
        result["outputs"].update({"所得税计算及政策依据": "calculation.json", "税表行次至原始财报来源": "lineage.json"})
        with (out / "workflow_report.html").open("x", encoding="utf-8") as stream:
            stream.write(html_report(result))
        result["outputs"]["业务流程结果报告"] = "workflow_report.html"
        tax_fill.dump_json(out / "workflow_result.json", result)
    return result


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    run = sub.add_parser("run", help="从原始利润表自动取累计数并生成 EIT 主表复核稿")
    run.add_argument("case", type=Path)
    run.add_argument("--out", required=True, type=Path)
    args = parser.parse_args(argv)
    try:
        result = run_case(args.case, args.out)
        print(json.dumps({"status": result["status"], "outcome": result["outcome"], "issues": len(result["issues"]),
                          "output_directory": str(args.out.absolute()), "report": str(args.out.absolute() / "workflow_report.html")}, ensure_ascii=False))
        return 2 if result["outcome"] == "BLOCKED" else 0
    except (tax_fill.Blocked, OSError, ValueError) as error:
        print(json.dumps({"status": "BLOCKED", "code": getattr(error, "code", "WORKFLOW_INPUT"), "message": str(error)}, ensure_ascii=False), file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
