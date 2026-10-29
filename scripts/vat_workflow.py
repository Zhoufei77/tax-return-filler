#!/usr/bin/env python3
"""Existing CSV/XLSX ledger -> confirmed 2026 VAT rules -> draft working papers.

This workflow never edits the ledger, logs in, files a return or pays tax.
The bundled forms reproduce official semantic rows; they are working papers,
not a claim of successful tax-system import.
"""
from __future__ import annotations

import argparse
from contextlib import redirect_stdout
from collections import defaultdict
import copy
import csv
from datetime import date, datetime
from decimal import Decimal
import html
import io
import json
from pathlib import Path
import re
import sys
import unicodedata

from tax_fill import Blocked, Source, amount, cell_address, digest, dump_json, input_path, serial, verify_hash
import tax_fill
import vat_rules

VERSION = "0.1.0"
FORMS_DIR = Path(__file__).resolve().parents[1] / "assets"
REQUIRED = ("id", "taxpayer_id", "tax_date", "category", "invoice_kind", "status", "net_amount", "applicable_rate")
ENUMS = {
    "category": {"goods_labor", "services"}, "invoice_kind": {"special", "ordinary", "unbilled"},
    "status": {"valid", "void", "red"},
}
HEADERS = {
    "id": {"业务行号", "业务流水号", "流水号", "记录编号", "业务编号", "发票号码", "数电发票号码", "全电发票号码"},
    "line_no": {"明细行号", "发票明细行号", "商品行号", "行号"},
    "taxpayer_id": {"纳税人识别号", "销方识别号", "销售方纳税人识别号", "销方纳税人识别号", "统一社会信用代码", "纳税人税号"},
    "tax_date": {"税款所属日期", "业务所属日期", "纳税义务发生日期", "增值税纳税义务日期", "纳税义务日期", "开票日期", "发票日期", "开具日期"},
    "category": {"业务类别", "业务分类", "交易类别", "销售类别", "业务类型"},
    "invoice_kind": {"票种", "发票票种", "发票种类", "发票类型"},
    "status": {"状态", "发票状态", "票据状态"},
    "net_amount": {"不含税金额", "不含税销售额", "销售额", "金额"},
    "applicable_rate": {"适用征收率", "适用税率", "税率", "征收率", "税率/征收率"},
}
BUILTIN_VALUES = {
    "category": {
        "货物": "goods_labor", "销售货物": "goods_labor", "货物销售": "goods_labor", "商品销售": "goods_labor",
        "加工修理修配": "goods_labor", "加工修理修配服务": "goods_labor", "货物及劳务": "goods_labor",
        "服务": "services", "其他服务": "services", "咨询服务": "services", "信息技术服务": "services", "软件服务": "services",
    },
    "invoice_kind": {
        "专票": "special", "增值税专用发票": "special", "增值税电子专用发票": "special",
        "电子发票(增值税专用发票)": "special", "数电发票(增值税专用发票)": "special", "数电专票": "special",
        "普票": "ordinary", "普通发票": "ordinary", "增值税普通发票": "ordinary", "增值税电子普通发票": "ordinary",
        "电子发票(普通发票)": "ordinary", "数电发票(普通发票)": "ordinary", "数电普票": "ordinary",
        "未开票": "unbilled", "未开具发票": "unbilled", "无票收入": "unbilled",
    },
    "status": {"正常": "valid", "有效": "valid", "正常发票": "valid", "作废": "void", "已作废": "void", "红字": "red", "红冲": "red", "已红冲": "red", "全额红冲": "red", "部分红冲": "red"},
}


def require(condition, code, message):
    if not condition:
        raise Blocked(code, message)


def norm(value):
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(value)))


def header_name(value):
    return re.sub(r"\((?:(?:金额)?单位[:：]?)?(?:人民币)?(?:元|千元|万元)\)$", "", norm(value))


def unit_declarations(value):
    """Explicit unit annotations remain binding even after a column override."""
    text = norm(value)
    found = [match.group(1) for match in re.finditer(r"(?:金额单位|单位)[:：]([^()【】\[\];；,，]+)", text)]
    for content in re.findall(r"[([【]([^()\[\]【】]+)[)\]】]", text):
        unit_content = re.sub(r"^(?:金额)?单位[:：]?", "", content)
        if unit_content.endswith("元"):
            found.append(unit_content)
    ending = re.search(r"(?:/|_|-)(万元|千元|元)$", text)
    if ending:
        found.append(ending.group(1))
    # A nonstandard annotation mentioning a larger currency unit must never
    # disappear merely because the Agent selected its column by letter.
    if not found:
        found.extend(re.findall(r"万元|千元", text))
    return list(dict.fromkeys(found))


def currency_declarations(value):
    text = norm(value)
    found = []
    codes = r"CNY|RMB|USD|EUR|HKD|JPY|GBP|AUD|CAD|SGD|CHF|NZD"
    found.extend(match.group(1) for match in re.finditer(rf"(?<![A-Za-z])({codes})(?![A-Za-z])", text, flags=re.I))
    found.extend(word for word in ("人民币", "美元", "美金", "欧元", "港币", "港元", "日元", "英镑", "澳元", "加元") if word in text)
    for content in re.findall(r"[([【]([^()\[\]【】]+)[)\]】]", text):
        if re.fullmatch(r"[A-Z]{3}", content):
            found.append(content)  # Unknown stated currency is not CNY by default.
        match = re.fullmatch(r"(?:币种|货币|CURRENCY)[:：](.+)", content, flags=re.I)
        if match:
            found.append(match.group(1))
    return list(dict.fromkeys(found))


def normalized_metadata(key, value):
    text = norm(value)
    if key == "currency" and text.upper() in ("CNY", "RMB"):
        return "CNY"
    if key == "currency" and text in ("人民币", "人民币元"):
        return "CNY"
    if key == "amount_unit" and text in ("人民币元", "CNY元", "RMB元"):
        return "元"
    return text


def meaningful(value):
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def evidence(item):
    return isinstance(item, dict) and "value" in item and isinstance(item.get("evidence"), str) and bool(item["evidence"].strip())


def json_value(value):
    return serial(value)


def date_value(raw, label):
    if isinstance(raw, datetime):
        return raw.date(), "xlsx_datetime_to_iso"
    if isinstance(raw, date):
        return raw, "xlsx_date_to_iso"
    require(isinstance(raw, str), "DATE_VALUE", f"{label} 必须是明确日期文本或 XLSX 日期")
    text = unicodedata.normalize("NFKC", raw).strip()
    match = re.fullmatch(r"(20[0-9]{2})[-/年]([0-9]{1,2})[-/月]([0-9]{1,2})日?", text)
    require(bool(match), "DATE_VALUE", f"{label} 日期格式不明确: {text}")
    try:
        return date(*map(int, match.groups())), "explicit_date_to_iso"
    except ValueError:
        raise Blocked("DATE_VALUE", f"{label} 日期无效: {text}")


def parse_amount(raw):
    converted = unicodedata.normalize("NFKC", raw).strip() if isinstance(raw, str) else raw
    value = amount(converted, "台账不含税金额")
    vat_rules.money(format(value, "f"), "net_amount")
    return format(value, "f"), "normalized_decimal_yuan"


def parse_rate(raw):
    if isinstance(raw, bool):
        raise Blocked("RATE_VALUE", "布尔值不能作为征收率")
    text = unicodedata.normalize("NFKC", str(raw)).strip()
    if text.endswith("%"):
        value = amount(text[:-1], "征收率") / Decimal(100)
    else:
        value = amount(raw if not isinstance(raw, str) else text, "征收率")
    require(value in (Decimal("0.01"), Decimal("0.03")), "RATE_VALUE", "仅支持明确1%或3%；票面免税/0不能自动改为适用征收率")
    return "0.01" if value == Decimal("0.01") else "0.03"


def resolve_header(reader, sheet, ledger):
    from openpyxl.utils.cell import get_column_letter
    rows, cols = reader.dimensions(sheet)
    explicit = ledger.get("columns", {})
    require(isinstance(explicit, dict) and set(explicit) <= set(HEADERS), "COLUMNS_SCHEMA", "columns 含未知字段或不是对象")
    requested = ledger.get("header_row")
    if requested is not None:
        require(type(requested) is int and 1 <= requested <= rows, "HEADER_ROW", "header_row 必须是实际表头行")
        candidates = [requested]
    else:
        candidates = []
        for row in range(1, min(rows, 25) + 1):
            names = set()
            for col in range(1, cols + 1):
                raw, kind = reader.raw(sheet, f"{get_column_letter(col)}{row}")
                if meaningful(raw) and kind not in ("f", "e"):
                    names.add(header_name(raw))
            recognized = {key for key, aliases in HEADERS.items() if names & aliases}
            if {"id", "tax_date", "net_amount"} <= recognized:
                candidates.append(row)
        require(len(candidates) == 1, "HEADER_AMBIGUOUS", "未识别唯一台账表头；请由Agent指定真实header_row和列映射，不要求用户改表")
    row = candidates[0]
    labels, positions = {}, defaultdict(list)
    for col in range(1, cols + 1):
        address = f"{get_column_letter(col)}{row}"
        raw, kind = reader.raw(sheet, address)
        if meaningful(raw):
            require(kind not in ("f", "e"), "HEADER_FORMULA", "表头含公式或错误，不能建立稳定映射")
            text = header_name(raw)
            labels[get_column_letter(col)] = str(raw)
            positions[text].append(get_column_letter(col))
    duplicates = [name for name, columns in positions.items() if len(columns) > 1]
    require(not duplicates, "DUPLICATE_HEADER", f"台账有同名列，不能任选: {duplicates}")
    mapping = {}
    for key, aliases in HEADERS.items():
        if key in explicit:
            selected = explicit[key]
            require(isinstance(selected, str) and selected.strip(), "COLUMN_MAPPING", "columns值须为实际表头文字或列字母")
            by_name = positions.get(header_name(selected), [])
            if by_name:
                candidates = by_name
            elif re.fullmatch(r"[A-Z]{1,3}", selected) and selected in labels:
                candidates = [selected]
            else:
                raise Blocked("COLUMN_MAPPING", f"明确指定的列不存在: {key}={selected}")
        else:
            candidates = [column for name, columns in positions.items() if name in aliases for column in columns]
            if key == "id":
                primary = [column for column in candidates if "发票" not in header_name(labels[column])]
                if primary:
                    candidates = primary
        require(len(candidates) <= 1, "COLUMN_AMBIGUOUS", f"{key} 同时匹配多个列，需显式选择: {candidates}")
        if candidates:
            column = candidates[0]
            mapping[key] = {"column": column, "header": labels[column], "header_cell": f"{column}{row}", "explicit": key in explicit}
    used = defaultdict(list)
    for key, entry in mapping.items():
        used[entry["column"]].append(key)
    require(all(len(keys) == 1 for keys in used.values()), "COLUMN_REUSED", "同一列不能同时绑定不同业务字段")
    require("id" in mapping and "tax_date" in mapping and "net_amount" in mapping, "COLUMN_REQUIRED", "业务键、税务日期、净额必须来自真实列，不能制造默认值")
    require(not re.search(r"价税合计|含税", norm(mapping["net_amount"]["header"]).replace("不含税", "")), "GROSS_AS_NET", "不能把价税合计或含税金额当成净额")
    return row, mapping, labels


def inspect_metadata(reader, sheet, header_row, mapping, ledger, case, context):
    from openpyxl.utils.cell import get_column_letter
    found = []
    rows, cols = reader.dimensions(sheet)
    for row in range(1, header_row):
        for col in range(1, cols + 1):
            cell = f"{get_column_letter(col)}{row}"
            raw, kind = reader.raw(sheet, cell)
            if not isinstance(raw, str) or kind in ("f", "e"):
                continue
            text = norm(raw)
            for match in re.finditer(r"(?:金额单位|单位)[:：]?(?:人民币)?(万元|千元|元)", text):
                found.append({"key": "amount_unit", "value": match.group(1), "cell": cell})
            for key, pattern in (("currency", r"币种[:：]([^;；,，]+)"), ("taxpayer_id", r"(?:纳税人识别号|销方识别号|统一社会信用代码)[:：]([^;；,，]+)")):
                match = re.search(pattern, text)
                if match:
                    found.append({"key": key, "value": match.group(1), "cell": cell})
            adjacent_roles = {"币种": "currency", "金额单位": "amount_unit", "单位": "amount_unit", "纳税人识别号": "taxpayer_id", "销方识别号": "taxpayer_id", "统一社会信用代码": "taxpayer_id"}
            role = adjacent_roles.get(text.rstrip(":"))
            if role and col < cols:
                next_cell = f"{get_column_letter(col + 1)}{row}"
                value, vk = reader.raw(sheet, next_cell)
                if meaningful(value):
                    require(vk not in ("f", "e"), "METADATA_FORMULA", "来源元数据含公式或错误")
                    found.append({"key": role, "value": str(value).strip(), "cell": next_cell})
    for entry in mapping.values():
        for unit in unit_declarations(entry["header"]):
            found.append({"key": "amount_unit", "value": unit, "cell": entry["header_cell"]})
    for currency in currency_declarations(mapping["net_amount"]["header"]):
        found.append({"key": "currency", "value": currency, "cell": mapping["net_amount"]["header_cell"]})
    declared = ledger.get("source_metadata", case.get("source_metadata", {}))
    require(isinstance(declared, dict), "SOURCE_METADATA", "source_metadata 必须为对象")
    for key in ("taxpayer_id", "currency", "amount_unit"):
        if key in declared:
            item = declared[key]
            value = item.get("value") if isinstance(item, dict) else item
            found.append({"key": key, "value": value, "origin": "explicit_source_metadata"})
    for entry in found:
        actual = normalized_metadata(entry["key"], entry["value"])
        require(actual == context[entry["key"]], "SOURCE_METADATA_CONFLICT", f"来源{entry['key']}={entry['value']}与已确认上下文冲突")
    return found


def default_value(ledger, key):
    defaults = ledger.get("defaults", {})
    require(isinstance(defaults, dict), "DEFAULT_SCHEMA", "defaults必须为对象")
    item = defaults.get(key)
    require(evidence(item), "DEFAULT_EVIDENCE", f"{key} 无真实列值时，默认值必须明确value及evidence，不能自动猜")
    require(meaningful(item["value"]), "DEFAULT_VALUE", f"{key} 默认值为空")
    return item["value"], {"kind": "confirmed_default", "value": json_value(item["value"]), "evidence": item["evidence"]}


def enum_value(key, raw, ledger):
    text = norm(raw)
    mappings = ledger.get("value_mappings", {})
    require(isinstance(mappings, dict), "VALUE_MAPPING", "value_mappings必须为对象")
    custom = mappings.get(key, {})
    require(isinstance(custom, dict), "VALUE_MAPPING", f"{key}值映射必须为对象")
    custom = {norm(k): v for k, v in custom.items()}
    if text in custom:
        result, conversion = custom[text], "explicit_value_mapping"
    elif text in ENUMS[key]:
        result, conversion = text, "already_normalized_enum"
    else:
        result, conversion = {norm(k): v for k, v in BUILTIN_VALUES[key].items()}.get(text), "exact_chinese_alias"
    require(result in ENUMS[key], "VALUE_UNRECOGNIZED", f"{key} 值无法确定映射: {raw}")
    return result, conversion


def extract_ledger(path, ledger, case, context):
    from openpyxl.utils.cell import get_column_letter
    sha = digest(path)
    reader = Source(path)
    # Source.raw revalidates dimensions. These read-only sources never change
    # in memory; cache dimensions once rather than rescanning all workbook cells
    # for every raw value in a large export.
    dimensions = {name: reader.dimensions(name) for name in reader.sheets}

    def cached_dimensions(name):
        require(name in dimensions, "SHEET_MISSING", f"来源工作表不存在: {name}")
        return dimensions[name]

    reader.dimensions = cached_dimensions
    try:
        sheet = ledger.get("sheet", "CSV" if path.suffix.lower() == ".csv" else reader.sheets[0] if len(reader.sheets) == 1 else None)
        require(sheet in reader.sheets, "SHEET_REQUIRED", "工作簿多表时需由Agent指定实际收入台账sheet")
        header_row, mapping, labels = resolve_header(reader, sheet, ledger)
        metadata = inspect_metadata(reader, sheet, header_row, mapping, ledger, case, context)
        ss = vat_rules.iso(ledger.get("source_period_start"), "source_period_start")
        se = vat_rules.iso(ledger.get("source_period_end"), "source_period_end")
        start, end = vat_rules.iso(context["period_start"], "period_start"), vat_rules.iso(context["period_end"], "period_end")
        require(ss <= start <= end <= se, "SOURCE_PERIOD", "明确源数据窗口必须覆盖目标期间")
        require(isinstance(ledger.get("period_evidence"), str) and ledger["period_evidence"].strip(), "SOURCE_PERIOD_EVIDENCE", "来源期间须有实际导出范围或资料确认的证据")
        date_header = header_name(mapping["tax_date"]["header"])
        invoice_date_header = date_header in ("开票日期", "发票日期", "开具日期")
        if invoice_date_header:
            confirmation = ledger.get("invoice_date_equals_tax_date")
            require(evidence(confirmation) and confirmation["value"] is True, "INVOICE_DATE_NOT_TAX_DATE", "只有开票日期时，需明确确认本批交易开票日期等于增值税义务日期，不能猜")
        result = {"path": str(path), "sha256": sha, "sheet": sheet, "header_row": header_row, "columns": mapping,
                  "metadata_evidence": metadata, "declared_context": {key: context[key] for key in ("taxpayer_id", "currency", "amount_unit")},
                  "declaration_evidence": case["confirmation_evidence"], "source_period_start": ss.isoformat(), "source_period_end": se.isoformat(),
                  "period_evidence": ledger["period_evidence"], "rows": [], "issues": [], "transactions": []}
        metadata_columns = {}
        for column, label in labels.items():
            text = header_name(label).upper()
            if text in ("币种", "币别", "币种代码", "货币代码", "CURRENCY"):
                metadata_columns[column] = "currency"
            elif text in ("金额单位", "货币单位", "金额计量单位"):
                metadata_columns[column] = "amount_unit"
            elif text == "单位":
                metadata_columns[column] = "possible_amount_unit"
        result["metadata_columns"] = metadata_columns
        rows, cols = reader.dimensions(sheet)
        seen = set()
        for row in range(header_row + 1, rows + 1):
            raw_row = {get_column_letter(col): json_value(reader.raw(sheet, f"{get_column_letter(col)}{row}")[0]) for col in range(1, cols + 1)}
            if not any(meaningful(value) for value in raw_row.values()):
                continue
            record = {"row": row, "raw_cells": raw_row, "fields": {}, "included": False}
            result["rows"].append(record)
            try:
                record["metadata_evidence"] = []
                for column, role in metadata_columns.items():
                    metadata_cell = f"{column}{row}"
                    metadata_raw, metadata_kind = reader.raw(sheet, metadata_cell)
                    if not meaningful(metadata_raw):
                        continue
                    require(metadata_kind not in ("f", "e"), "METADATA_FORMULA", f"行内币种/单位含公式或错误: {metadata_cell}")
                    if role == "possible_amount_unit":
                        # Invoice quantity '单位=件' is not a currency unit.
                        if norm(metadata_raw) not in ("元", "千元", "万元", "人民币元", "CNY元", "RMB元"):
                            continue
                        role = "amount_unit"
                    value = normalized_metadata(role, metadata_raw)
                    record["metadata_evidence"].append({"key": role, "raw": json_value(metadata_raw), "value": value, "cell": metadata_cell})
                    require(value == context[role], "SOURCE_METADATA_CONFLICT", f"第{row}行来源{role}={metadata_raw}与已确认上下文冲突")
                normalized = {}
                for key in REQUIRED:
                    entry = mapping.get(key)
                    if entry:
                        address = f"{entry['column']}{row}"
                        raw, kind = reader.raw(sheet, address)
                        require(kind not in ("f", "e"), "FORMULA_SOURCE", f"源字段含公式或错误，不能读缓存: {sheet}!{address}")
                    else:
                        raw, address, kind = None, None, None
                    if not meaningful(raw):
                        require(key not in ("id", "tax_date", "net_amount"), "REQUIRED_SOURCE_VALUE", f"{key} 缺少源值: 第{row}行")
                        raw, origin = default_value(ledger, key)
                    else:
                        origin = {"kind": "source_cell", "cell": address, "header_cell": entry["header_cell"], "header": entry["header"], "raw_type": kind}
                    conversion = "text_preserved"
                    if key == "id":
                        invoice_number = "发票" in header_name(entry["header"])
                        require(isinstance(raw, str) or (type(raw) is int and not invoice_number), "IDENTIFIER_PRECISION", "发票号码必须为文本；业务流水号只支持明确文本或安全整数")
                        require(not isinstance(raw, int) or len(str(abs(raw))) <= 15, "IDENTIFIER_PRECISION", "超过15位的数值流水号可能已经丢精度，请使用保留原文的导出")
                        value = str(raw).strip()
                        require(value and not value.startswith(("=", "+", "-", "@")), "IDENTIFIER_VALUE", "交易键为空或可能是公式")
                        if "line_no" in mapping:
                            line_entry = mapping["line_no"]
                            line_cell = f"{line_entry['column']}{row}"
                            line_raw = reader.value(sheet, line_cell)
                            require(isinstance(line_raw, str) or type(line_raw) is int, "LINE_KEY", "明细行号必须为文本或整数")
                            value = json.dumps([value, str(line_raw).strip()], ensure_ascii=False, separators=(",", ":"))
                            origin["line_no"] = {"cell": line_cell, "value": json_value(line_raw)}
                            conversion = "composite_invoice_or_business_and_line_key"
                    elif key == "taxpayer_id":
                        require(isinstance(raw, str), "TAXPAYER_PRECISION", "源税号必须为文本，数值型税号可能失真")
                        value = raw.strip()
                        require(value == context["taxpayer_id"], "TAXPAYER_MIXED", f"台账混入其他主体: 第{row}行")
                    elif key == "tax_date":
                        dt, conversion = date_value(raw, f"{sheet}第{row}行")
                        require(ss <= dt <= se, "SOURCE_DATE_WINDOW", f"第{row}行日期超出已声明来源窗口")
                        value = dt.isoformat()
                        if invoice_date_header:
                            origin["tax_date_confirmation"] = ledger["invoice_date_equals_tax_date"]
                    elif key == "net_amount":
                        value, conversion = parse_amount(raw)
                    elif key == "applicable_rate":
                        try:
                            value = parse_rate(raw)
                            conversion = "explicit_fraction_or_percent_to_rate"
                        except Blocked:
                            require(norm(raw) in ("免税", "0", "0.0", "0.00", "0%", "0.0%", "0.00%"), "RATE_VALUE", "不明确的征收率不能默认替代")
                            replacement, default_origin = default_value(ledger, key)
                            value = parse_rate(replacement)
                            origin["applicable_rate_confirmation"] = default_origin
                            conversion = "invoice_zero_or_exempt_to_separately_confirmed_applicable_rate"
                    else:
                        value, conversion = enum_value(key, raw, ledger)
                    normalized[key] = value
                    record["fields"][key] = {"raw": json_value(raw), "normalized": value, "conversion": conversion, **origin}
                require(normalized["id"] not in seen, "DUPLICATE_RECORD", f"源台账交易键重复，不能自动去重: 第{row}行")
                seen.add(normalized["id"])
                source = {"path": str(path), "sha256": sha, "sheet": sheet, "row": row,
                          "cells": {key: entry.get("cell") for key, entry in record["fields"].items()}, "lineage_row": len(result["rows"]) - 1}
                normalized["source"] = source
                dt = date.fromisoformat(normalized["tax_date"])
                take = start <= dt <= end
                record.update({"id": normalized["id"], "included": take, "reason": "目标所属期内；后续由规则处理有效/作废" if take else "在已确认源窗口内、目标期间外，显式排除，未改期", "normalized": normalized})
                if take:
                    result["transactions"].append(normalized)
            except (Blocked, vat_rules.InputError) as exc:
                code = exc.code if isinstance(exc, Blocked) else "VALUE_INVALID"
                result["issues"].append({"code": code, "message": str(exc), "row": row, "sheet": sheet})
                record["reason"] = "源行校验失败，不能进入计算"
        require(digest(path) == sha, "HASH_CHANGED", "读取过程中源台账变化，请重新扫描")
        return result
    finally:
        if reader.workbook:
            reader.workbook.close()


def check_target(workbook, target):
    require(isinstance(target, dict), "FORM_TARGET", "套表目标不是对象")
    sheet, cell = target.get("sheet"), cell_address(target.get("cell"))
    require(sheet in workbook.sheetnames, "FORM_TARGET", f"套表工作表不存在: {sheet}")
    anchor = target.get("anchor", {})
    anchor_cell = cell_address(anchor.get("cell"))
    require(workbook[sheet][anchor_cell].value == anchor.get("text"), "FORM_ANCHOR", f"套表标签锚点不匹配: {sheet}!{anchor_cell}")
    from openpyxl.cell.cell import MergedCell
    c = workbook[sheet][cell]
    require(not isinstance(c, MergedCell) and c.data_type not in ("f", "e"), "FORM_TARGET", f"目标不能是公式、错误或合并非首格: {sheet}!{cell}")
    return c


def prepare_fill(case, calculation, output, template, meta, metadata_path, input_records):
    context = case["context"]
    reader = Source(template, template=True)
    wb = reader.workbook
    try:
        identity = meta.get("identity_fields")
        require(isinstance(identity, list) and len(identity) == 12 and len(wb.sheetnames) == 3, "FORM_IDENTITY", "内置三张套表必须各有税号、名称和起止日期，共12项")
        identity_roles = defaultdict(set)
        for entry in identity:
            require(isinstance(entry, dict) and isinstance(entry.get("target"), dict), "FORM_IDENTITY", "身份目标结构无效")
            identity_roles[entry["target"].get("sheet")].add(entry.get("context_key"))
        require(set(identity_roles) == set(wb.sheetnames) and all(roles == {"taxpayer_id", "taxpayer_name", "period_start", "period_end"} for roles in identity_roles.values()), "FORM_IDENTITY", "每张套表须绑定完整一致的主体和所属期")
        values = {"current": [], "ytd": []}
        fields, ratios = [], []
        targets = set()

        def append_field(fid, label, value, target, *, kind="money", basis="current", path=None):
            key = (target.get("sheet"), cell_address(target.get("cell")))
            require(key not in targets, "FORM_TARGET_DUPLICATE", f"套表目标重复: {key}")
            targets.add(key)
            cell = check_target(wb, target)
            require(cell.value is None or cell.value == "" or (str(cell.value) == str(value)), "FORM_OCCUPIED", f"内置填写目标已有不同值: {key}")
            bucket = values[basis]
            bucket.append([fid, value, kind, path or label])
            fields.append({"id": fid, "label": label, "kind": kind, "basis": basis, "required": True,
                           "target": {key: value for key, value in target.items() if key in ("sheet", "cell", "anchor")},
                           "expression": {"op": "ref", "source": basis, "sheet": "CSV", "cell": f"B{len(bucket)+1}"}})

        for index, entry in enumerate(identity):
            key = entry.get("context_key")
            require(key in context and isinstance(context[key], str), "FORM_IDENTITY", f"上下文缺少文本身份/期间值: {key}")
            append_field(f"identity_{index}", entry.get("label", key), context[key], entry["target"], kind="text", path=f"context.{key}")
        line_targets = meta.get("line_targets")
        require(isinstance(line_targets, dict) and set(line_targets) == {str(row) for row in range(1, 26)}, "FORM_LINES", "内置主表必须覆盖1—25语义栏次")
        not_applicable = defaultdict(set)
        for entry in meta.get("not_applicable_targets", []):
            rid, key = str(entry.get("row")), entry.get("key")
            require(key not in not_applicable[rid], "FORM_GRAY_DUPLICATE", "灰格声明重复")
            require(key in calculation["lines"].get(rid, {}), "FORM_GRAY_MISSING", "灰格对应计算值缺失，不能默默跳过")
            require(Decimal(calculation["lines"][rid][key]) == Decimal(entry.get("expected", "0.00")) == 0, "FORM_GRAY_NONZERO", f"主表{rid}/{key}为不可填灰格，但规则值非零，不能漏掉")
            not_applicable[rid].add(key)
        for rid, targets_by_kind in line_targets.items():
            expected = set(calculation["lines"].get(rid, {})) - not_applicable[rid]
            require(isinstance(targets_by_kind, dict) and set(targets_by_kind) == expected, "FORM_LINES", f"主表{rid}可填列与已核实计算结果不一致")
            for key, target in targets_by_kind.items():
                basis = "ytd" if key.startswith("ytd_") else "current"
                append_field(f"vat_{rid}_{key}", f"增值税主表第{rid}栏 {key}", calculation["lines"][rid][key], target, basis=basis, path=f"calculation.lines.{rid}.{key}")
        for name, targets_by_column in meta.get("surcharge_targets", {}).items():
            detail = calculation["surcharge_details"].get(name)
            require(detail is not None, "FORM_SURCHARGE", f"缺少已确认附加明细: {name}")
            require(set(targets_by_column) == {str(col) for col in range(1, 10)}, "FORM_SURCHARGE", "附加表必须覆盖1—9列")
            for column, target in targets_by_column.items():
                value = detail["columns"][column]
                unit = target.get("unit", "text" if column == "4" else "rate" if column in ("2", "6") else "money")
                if unit == "rate":
                    key = (target.get("sheet"), cell_address(target.get("cell")))
                    require(key not in targets, "FORM_TARGET_DUPLICATE", "比率目标与其他字段重复")
                    targets.add(key)
                    c = check_target(wb, target)
                    before = serial(c.value)
                    ratio = Decimal(value)
                    if c.value is not None and not (name == "urban" and column == "2"):
                        require(amount(c.value, "模板固定比率") == ratio, "FORM_FIXED_RATE", "模板已有固定3%/2%/50%与核实结果不一致")
                    c.value = float(ratio)
                    c.number_format = "0.00%"
                    ratios.append({"sheet": target["sheet"], "cell": target["cell"], "before": before, "ratio": str(ratio), "format": "0.00%", "calculation_path": f"surcharge_details.{name}.columns.{column}", "rule_sources": detail["rule_sources"]})
                elif unit == "text":
                    if meaningful(value):
                        append_field(f"surcharge_{name}_{column}", f"{name}附加第{column}列", value, target, kind="text", path=f"surcharge_details.{name}.columns.{column}")
                else:
                    require(unit == "money", "FORM_UNIT", "附加目标单位未知")
                    append_field(f"surcharge_{name}_{column}", f"{name}附加第{column}列", value, target, path=f"surcharge_details.{name}.columns.{column}")
        require(set(meta.get("surcharge_targets", {})) == {"urban", "education", "local_education"}, "FORM_SURCHARGE", "内置套表缺少附加税费行")
        relief = calculation["relief_details"]
        if relief["required"]:
            targets_by_column = meta.get("relief_targets", {})
            require(set(targets_by_column) == {"code", "1", "2", "3", "4", "5"}, "FORM_RELIEF", "内置减免明细缺少代码或1—5列")
            append_field("relief_code", "1%减征性质代码", relief["reduction_1pct"]["code"], targets_by_column["code"], kind="text", path="relief_details.reduction_1pct.code")
            for column in ("1", "2", "3", "4", "5"):
                append_field(f"relief_{column}", f"1%减征第{column}列", relief["reduction_1pct"]["columns"][column], targets_by_column[column], path=f"relief_details.reduction_1pct.columns.{column}")
        derived = output / "derived-template.xlsx"
        wb.save(derived)
    finally:
        wb.close()
    sources = []
    for basis, rows in values.items():
        if not rows:
            continue
        path = output / f"calculated-{basis}.csv"
        with path.open("x", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["字段", "已计算值", "类型", "计算来源"])
            writer.writerows(rows)
        sources.append({"id": basis, "path": str(path), "sha256": digest(path), "metadata_confirmed": True,
                        "context": {"taxpayer_id": context["taxpayer_id"], "currency": "CNY", "amount_unit": "元",
                                    "period_start": context["period_start"] if basis == "current" else f"{date.fromisoformat(context['period_start']).year}-01-01", "period_end": context["period_end"]}})
    # Ratio cells are already written as typed fractions in the derived template.
    # They are excluded from core money targets; all monetary values use CSV ref.
    job = {"schema_version": 1, "context": context, "confirmations": {"context_confirmed": True, "materials_complete": True},
           "sources": sources, "template": {"path": str(derived), "sha256": digest(derived), "context_bindings": meta["context_bindings"]},
           "fields": fields, "mandatory_targets": [field["target"] for field in fields], "checks": []}
    job_path = output / "fill-job.json"
    dump_json(job_path, job)
    return job_path, {"base_template_sha256": digest(template), "metadata_sha256": digest(metadata_path), "derived_template_sha256": digest(derived), "ratio_writes": ratios, "derived_value_sources": sources, "confirmation_evidence": case["confirmation_evidence"], "input_records": input_records}


def verify_filled(output, preparation):
    """Read back serialized cells, ratios and preserved formulas before delivery."""
    audit = json.loads((output / "fill-result" / "audit.json").read_text(encoding="utf-8"))
    require(audit.get("outcome") == "GENERATED_FOR_REVIEW" and not audit.get("issues"), "FILL_AUDIT", "core审计未确认机械填写完成")
    reader = Source(output / "fill-result" / "filled.xlsx", template=True)
    base = Source(output / "derived-template.xlsx", template=True)
    try:
        require(reader.sheets == base.sheets, "OUTPUT_SHEETS", "填写稿工作表结构变化")
        for field in audit["fields"]:
            target = field["target"]
            raw = reader.value(target["sheet"], target["cell"])
            if field["kind"] == "money":
                require(amount(raw, "填写稿回读") == Decimal(field["value"]), "OUTPUT_VALUE", f"金额序列化后不符: {target}")
            else:
                require(str(raw) == field["value"], "OUTPUT_VALUE", f"身份/期间/代码序列化后不符: {target}")
        for ratio in preparation["ratio_writes"]:
            raw = reader.value(ratio["sheet"], ratio["cell"])
            require(amount(raw, "比例回读") == Decimal(ratio["ratio"]), "OUTPUT_RATIO", "比例保存后变化")
            require("%" in reader.workbook[ratio["sheet"]][ratio["cell"]].number_format, "OUTPUT_RATIO_FORMAT", "比例未保留百分数显示格式")
        formulas = 0
        for ws in base.workbook:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        require(reader.workbook[ws.title][cell.coordinate].value == cell.value, "OUTPUT_FORMULA", "原有公式在填写后变化")
                        formulas += 1
        return {"read_back_fields": len(audit["fields"]), "read_back_ratio_cells": len(preparation["ratio_writes"]), "preserved_formula_count": formulas, "all_expected_values_match": True}
    finally:
        reader.workbook.close()
        base.workbook.close()


def render_review(report, output):
    esc = lambda value: html.escape(str(value), quote=True)
    issues = "".join(f"<li><strong>{esc(item.get('code'))}</strong> {esc(item.get('message'))}</li>" for item in report["issues"])
    if not issues:
        issues = "<li>本次约定范围的自动检查通过，填写稿仍需经办会计复核。</li>"
    draft = output / "fill-result" / "filled.xlsx"
    link = '<p><a href="fill-result/filled.xlsx">打开发票与台账计算后的填写稿</a> · <a href="fill-result/review.html">逐字段填写复核</a></p>' if report["outcome"] == "GENERATED_FOR_REVIEW" and draft.exists() else "<p>本次有阻断项，没有可交付填写稿。</p>"
    content = f'<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>增值税填表复核</title><style>body{{font:16px/1.6 system-ui;max-width:1000px;margin:32px auto;padding:0 20px}}aside{{background:#fff3cf;padding:16px}}li{{margin:10px 0}}</style><h1>增值税填表复核</h1><aside>{esc(report["status"])} · {esc(report["outcome"])}<br>这是官方语义行次重建的申报工作底稿，未提交申报、未缴税。没有税局导入实测。</aside>{link}<h2>检查结果</h2><ul>{issues}</ul><p><a href="calculation.json">税额及规则计算记录</a> · <a href="lineage.json">原始台账来源、行列与转换</a> · <a href="workflow.json">本次执行记录</a></p></html>'
    with (output / "review.html").open("x", encoding="utf-8") as handle:
        handle.write(content)


def run(case_path, out, *, forms_dir=None):
    out = Path(out).expanduser().absolute()
    require(not out.exists() and not out.is_symlink(), "OUTPUT_EXISTS", "--out必须是尚不存在的新目录")
    case_path = input_path(Path.cwd(), str(case_path))
    require(out.resolve() != case_path and out.resolve() not in case_path.parents, "OUTPUT_INPUT", "输出不能包含或覆盖输入任务")
    report = {"schema_version": 1, "workflow_version": VERSION, "status": vat_rules.STATUS, "outcome": "BLOCKED", "issues": [], "inputs": [], "filing_performed": False, "human_review_required": True}
    report["implementation_sha256"] = {"workflow": digest(Path(__file__).resolve()), "vat_rules": digest(Path(vat_rules.__file__).resolve()), "tax_fill": digest(Path(tax_fill.__file__).resolve())}
    lineage, calculation = {}, {"status": vat_rules.STATUS, "outcome": "BLOCKED", "issues": []}
    created = False
    inputs = []
    try:
        require(case_path.suffix.lower() == ".json", "CASE_FORMAT", "case须为JSON")
        case_hash = digest(case_path)
        case = json.loads(case_path.read_text(encoding="utf-8-sig"))
        require(isinstance(case, dict), "CASE_FORMAT", "case顶层必须是对象")
        context = case.get("context", {})
        require(isinstance(context, dict), "CONTEXT", "context必须为对象")
        for key in ("taxpayer_id", "taxpayer_name", "period_start", "period_end", "period_kind", "form_name", "form_version", "region", "currency", "amount_unit"):
            require(isinstance(context.get(key), str) and context[key].strip(), "CONTEXT", f"缺少明确context.{key}")
        require(context["currency"] == "CNY" and context["amount_unit"] == "元", "UNIT_CURRENCY", "当前台账自动计算仅支持CNY元；不自动换汇或把万元当元")
        require(isinstance(case.get("confirmation_evidence"), str) and case["confirmation_evidence"].strip(), "CONFIRMATION_EVIDENCE", "须保存真实用户确认记录，不能由Agent自动猜为确认")
        raw_root = Path(case.get("input_dir", "")).expanduser()
        require(bool(case.get("input_dir")), "INPUT_DIRECTORY", "须明确本次input_dir")
        root = (raw_root if raw_root.is_absolute() else case_path.parent / raw_root).resolve()
        require(root.is_dir(), "INPUT_DIRECTORY", "input_dir不存在")
        ledger = case.get("ledger")
        require(isinstance(ledger, dict), "LEDGER", "ledger须为对象")
        path = input_path(root, ledger.get("path"))
        require(path.is_relative_to(root), "INPUT_SCOPE", "ledger不在用户指定input_dir中")
        require(out.resolve() != path and out.resolve() not in path.parents, "OUTPUT_INPUT", "输出不能包含或覆盖台账")
        inputs = [{"role": "case", "path": str(case_path), "sha256": case_hash}, {"role": "ledger", "path": str(path), "sha256": digest(path)}]
        report["inputs"] = inputs
        lineage = extract_ledger(path, ledger, case, context)
        require(not lineage["issues"], "LEDGER_INVALID", "源台账存在阻断项，详见lineage.json逐行位置")
        profile = copy.deepcopy(case.get("profile", {}))
        require(isinstance(profile, dict), "PROFILE", "profile须为对象")
        for key in ("taxpayer_id", "period_start", "period_end", "period_kind", "form_version", "region"):
            require(key not in profile or profile[key] == context[key], "PROFILE_CONTEXT", f"profile.{key}与context冲突")
            profile[key] = context[key]
        import fact_sources
        profile, fact_manifests, fact_trace = fact_sources.resolve_sources(profile, root, context)
        report["implementation_sha256"]["fact_sources"] = digest(Path(fact_sources.__file__).resolve())
        lineage["profile_fact_sources"] = fact_trace
        lineage["profile_fact_source_manifests"] = fact_manifests
        known_paths = {item["path"]: item["sha256"] for item in inputs}
        for item in fact_manifests:
            require(isinstance(item, dict) and isinstance(item.get("path"), str), "FACT_MANIFEST", "事实来源manifest不完整")
            if item["path"] in known_paths:
                require(known_paths[item["path"]] == item.get("sha256"), "HASH_CHANGED", "同一来源在台账与事实取数时内容变化")
            else:
                inputs.append({"role": "fact_source", **item})
                known_paths[item["path"]] = item["sha256"]
        calculation = vat_rules.calculate(lineage["transactions"], profile)
        lineage["calculation_file"] = "calculation.json"
        lineage["rule_sources"] = calculation["rule_sources"]
        require(calculation["outcome"] == "GENERATED_FOR_REVIEW", "CALCULATION_BLOCKED", "税额规则尚有缺项或不支持业务，详见calculation.json")
        forms = Path(forms_dir) if forms_dir else FORMS_DIR
        metadata_path, template = forms / "VAT_SMALL_2026.json", forms / "VAT_SMALL_2026.xlsx"
        require(metadata_path.is_file() and template.is_file(), "FORM_ASSET_MISSING", "缺少配套VAT_SMALL_2026.xlsx/json，不能冒充完成套表")
        meta = json.loads(metadata_path.read_text(encoding="utf-8"))
        require(meta.get("form_version") == context["form_version"], "FORM_VERSION", "内置套表版本与实际任务不符")
        verify_hash(template, meta.get("template_sha256"))
        inputs.extend([{"role": "template", "path": str(template.resolve()), "sha256": digest(template)}, {"role": "form_metadata", "path": str(metadata_path.resolve()), "sha256": digest(metadata_path)}])
        for item in inputs:
            verify_hash(Path(item["path"]), item["sha256"])
        out.mkdir(parents=True, exist_ok=False)
        created = True
        job_path, preparation = prepare_fill(case, calculation, out, template, meta, metadata_path, inputs)
        lineage["preparation"] = preparation
        with redirect_stdout(io.StringIO()):
            fill_exit = tax_fill.run(job_path, out / "fill-result")
        require(fill_exit == 0, "FILL_BLOCKED", "模板填写机械检查未通过，见fill-result/audit.json")
        lineage["preparation"]["output_verification"] = verify_filled(out, preparation)
        for item in inputs:
            verify_hash(Path(item["path"]), item["sha256"])
        report["outcome"] = "GENERATED_FOR_REVIEW"
        report["filled_workbook"] = str(out / "fill-result" / "filled.xlsx")
    except Blocked as exc:
        report["issues"].append({"code": exc.code, "message": exc.message})
    except (ValueError, TypeError, KeyError, AttributeError, OSError) as exc:
        report["issues"].append({"code": "WORKFLOW_INPUT", "message": f"{type(exc).__name__}: {exc}"})
    except Exception as exc:
        report["issues"].append({"code": "WORKFLOW_ERROR", "message": f"{type(exc).__name__}: {exc}"})
    finally:
        if not created:
            out.mkdir(parents=True, exist_ok=False)
        if report["issues"]:
            draft = out / "fill-result" / "filled.xlsx"
            if draft.exists():
                draft.unlink()
        if lineage.get("issues"):
            report["issues"].extend(lineage["issues"])
        if calculation.get("issues"):
            report["issues"].extend(calculation["issues"])
        dump_json(out / "lineage.json", lineage)
        dump_json(out / "calculation.json", calculation)
        dump_json(out / "workflow.json", report)
        render_review(report, out)
    print(json.dumps({"status": report["status"], "outcome": report["outcome"], "issues": len(report["issues"]), "output_directory": str(out)}, ensure_ascii=False))
    return 0 if report["outcome"] == "GENERATED_FOR_REVIEW" else 2


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    rp = sub.add_parser("run", help="实际台账取数、计算并生成增值税套表工作底稿")
    rp.add_argument("case", type=Path)
    rp.add_argument("--out", type=Path, required=True)
    args = parser.parse_args(argv)
    try:
        return run(args.case, args.out)
    except (Blocked, OSError, ValueError) as exc:
        print(json.dumps({"status": "BLOCKED", "message": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
