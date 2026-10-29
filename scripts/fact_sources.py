#!/usr/bin/env python3
"""Resolve typed tax facts from real, read-only CSV/XLSX source cells.

No network calls, workbook writes, formula evaluation, currency conversion,
or invented human confirmations are performed. Exceptions use tax_fill.Blocked.
"""
from __future__ import annotations

from copy import deepcopy
from datetime import date, datetime
from decimal import Decimal, localcontext
from pathlib import Path
import re

from statement_adapter import norm, parse_money
from tax_fill import Blocked, Source, UNITS, cell_address, digest, input_path, iso_date, serial, verify_hash

VERSION = "0.1.0"
IDENTITY_KEYS = {"taxpayer_id", "taxpayer_name", "period_end"}
UNIT_LABELS = {"单位", "金额单位", "计量单位", "货币单位"}
CURRENCY_LABELS = {"币种", "货币", "Currency", "currency", "记账本位币"}
CURRENCY_NAMES = {
    "人民币": "CNY", "CNY": "CNY", "RMB": "CNY", "美元": "USD", "USD": "USD",
    "欧元": "EUR", "EUR": "EUR", "港币": "HKD", "港元": "HKD", "HKD": "HKD",
    "日元": "JPY", "JPY": "JPY", "英镑": "GBP", "GBP": "GBP", "AUD": "AUD",
    "澳元": "AUD", "CAD": "CAD", "加元": "CAD", "CHF": "CHF", "SGD": "SGD", "新加坡元": "SGD",
}
CUR_PATTERN = "|".join(sorted(CURRENCY_NAMES, key=len, reverse=True))
UNIT_PATTERN = re.compile(r"(?:^|[;；,，(])(?:金额单位|货币单位|计量单位|单位)[:：]?(?:人民币|CNY|RMB)?(万元|千元|元)(?:人民币)?(?=$|[;；,，()])", re.I)
UNIT_PAREN_PATTERN = re.compile(r"(?:^|[;；,，(])(?:金额单位|货币单位|计量单位|单位)\((?:人民币|CNY|RMB)?(万元|千元|元)(?:人民币)?\)", re.I)
INLINE_UNIT_DECLARATION = re.compile(r"(?:^|[;；,，(])(?:金额单位|货币单位|计量单位|单位)(?:[:：]|\()([^;；,，()]+)", re.I)
CURRENCY_PATTERN = re.compile(r"(?:币种|记账本位币|Currency|货币)[:：]?(" + CUR_PATTERN + r")", re.I)
FOREIGN_UNIT_PATTERN = re.compile(r"(?:^|[;；,，(])(?:金额单位|货币单位|计量单位|单位)[:：]?(?:千|万)?(" + CUR_PATTERN + r")", re.I)


def need(condition, code, message):
    if not condition:
        raise Blocked(code, message)


def json_value(value):
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, dict):
        return {str(k): json_value(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_value(v) for v in value]
    return value


class Resolver:
    def __init__(self, input_dir, context):
        self.original_root = Path(input_dir).expanduser().absolute()
        need(not self.original_root.is_symlink(), "FACT_SOURCE_SYMLINK", "资料目录本身不能是符号链接")
        self.root = self.original_root.resolve()
        need(self.root.is_dir(), "FACT_INPUT_DIRECTORY", "事实来源资料目录不存在")
        need(isinstance(context, dict), "FACT_CONTEXT", "context 必须为对象")
        self.context = context
        self.readers, self.manifests, self.entries_cache, self.trace = {}, {}, {}, []

    def path(self, value):
        need(isinstance(value, str) and value.strip(), "FACT_SOURCE_PATH", "source.path 缺失")
        raw = Path(value).expanduser()
        candidate = raw if raw.is_absolute() else self.original_root / raw
        resolved = candidate.resolve()
        need(resolved.is_relative_to(self.root), "FACT_SOURCE_OUTSIDE_INPUT", "事实来源必须位于本次 input_dir 内，不能读取目录外资料")
        # Inspect lexical components below the caller's root as well as the
        # canonical root; resolving first must not hide an internal symlink.
        lexical_root = next((root for root in (self.original_root, self.root) if candidate.is_relative_to(root)), None)
        need(lexical_root is not None, "FACT_SOURCE_PATH_ALIAS", "绝对来源路径须直接位于 input_dir 或其规范路径下，不能通过目录外别名绕入")
        part_path = lexical_root
        for part in candidate.relative_to(lexical_root).parts:
            part_path = part_path / part
            need(not part_path.is_symlink(), "FACT_SOURCE_SYMLINK", "事实来源路径含符号链接，不读取其目标")
        path = input_path(self.root, str(resolved))
        need(path.suffix.lower() in (".xlsx", ".csv"), "FACT_SOURCE_FORMAT", "事实来源只支持 CSV/XLSX，不自动运行宏或转换其他格式")
        return path

    def reader(self, spec):
        path = self.path(spec.get("path"))
        if path not in self.readers:
            actual = digest(path)
            if "sha256" in spec:
                verify_hash(path, spec["sha256"])
            reader = Source(path)
            self.readers[path] = reader
            self.manifests[path] = {"path": str(path), "sha256": actual, "size": path.stat().st_size, "source_kind": "fact_source"}
            need(digest(path) == actual, "HASH_CHANGED", "事实来源在读取期间变化，必须重新扫描")
        elif "sha256" in spec:
            verify_hash(path, spec["sha256"])
        return path, self.readers[path], self.manifests[path]["sha256"]

    def entries(self, path, reader, sheet):
        key = (path, sheet)
        if key in self.entries_cache:
            return self.entries_cache[key]
        reader.dimensions(sheet)
        entries = []
        if reader.workbook:
            for row in reader.workbook[sheet].iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.data_type not in ("f", "e"):
                        entries.append((cell.row, cell.column, cell.coordinate, cell.value))
        else:
            from openpyxl.utils.cell import get_column_letter
            for row_index, row in enumerate(reader.rows, 1):
                for col_index, value in enumerate(row, 1):
                    if isinstance(value, str) and value and not value.startswith("="):
                        entries.append((row_index, col_index, f"{get_column_letter(col_index)}{row_index}", value))
        self.entries_cache[key] = entries
        return entries

    def metadata(self, path, reader, sheet, target, kind, declared_unit):
        from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
        target_row, target_col = coordinate_to_tuple(target)
        units, currencies, unknown_units = [], [], []
        unit_seen, currency_seen = set(), set()

        def unit(cell, raw, name, scope):
            marker = (cell, name, scope)
            if marker not in unit_seen:
                unit_seen.add(marker)
                units.append({"cell": cell, "raw": raw, "unit": name, "scope": scope})

        def currency(cell, raw, name):
            normalized = CURRENCY_NAMES.get(name.upper(), CURRENCY_NAMES.get(name))
            if normalized and (cell, normalized) not in currency_seen:
                currency_seen.add((cell, normalized))
                currencies.append({"cell": cell, "raw": raw, "currency": normalized})

        for row, col, cell, raw in self.entries(path, reader, sheet):
            value = norm(raw)
            for match in CURRENCY_PATTERN.finditer(value):
                currency(cell, raw, match[1])
            foreign = FOREIGN_UNIT_PATTERN.search(value)
            if foreign:
                currency(cell, raw, foreign[1])
            if row <= 20 and value.upper() in CURRENCY_NAMES:
                currency(cell, raw, value.upper())
            if row <= 20 and value in CURRENCY_NAMES:
                currency(cell, raw, value)
            for match in UNIT_PATTERN.finditer(value):
                unit(cell, raw, match[1], "sheet_declaration")
            for match in UNIT_PAREN_PATTERN.finditer(value):
                unit(cell, raw, match[1], "sheet_declaration")
            for declaration in INLINE_UNIT_DECLARATION.finditer(value):
                payload = declaration[1]
                known = re.fullmatch(r"(?:人民币|CNY|RMB)?(万元|千元|元)(?:人民币)?", payload, re.I)
                foreign_name = re.sub(r"^(?:千|万)", "", payload)
                if not known and foreign_name.upper() not in CURRENCY_NAMES and foreign_name not in CURRENCY_NAMES:
                    unknown_units.append({"cell": cell, "raw": raw, "declaration": payload})
            label = value.rstrip(":：")
            if label in UNIT_LABELS or label in CURRENCY_LABELS:
                neighbor_cell = f"{get_column_letter(col + 1)}{row}"
                adjacent, adjacent_kind = reader.raw(sheet, neighbor_cell)
                if adjacent is not None and adjacent != "":
                    need(adjacent_kind not in ("f", "e"), "FACT_METADATA_FORMULA", f"单位/币种元数据 {sheet}!{neighbor_cell} 为公式或错误，不能猜读缓存")
                    adjacent_text = norm(adjacent)
                    if label in CURRENCY_LABELS:
                        currency(neighbor_cell, str(adjacent), adjacent_text.upper() if adjacent_text.upper() in CURRENCY_NAMES else adjacent_text)
                    else:
                        matched = re.fullmatch(r"(?:人民币|CNY|RMB)?(万元|千元|元)(?:人民币)?", adjacent_text, re.I)
                        if matched:
                            unit(neighbor_cell, str(adjacent), matched[1], "sheet_declaration")
                        elif kind == "money" and adjacent_text not in ("",):
                            foreign_name = re.sub(r"^(?:千|万)", "", adjacent_text)
                            currency(neighbor_cell, str(adjacent), foreign_name.upper() if foreign_name.upper() in CURRENCY_NAMES else foreign_name)
                            if not any(c["currency"] != "CNY" for c in currencies):
                                raise Blocked("FACT_UNIT_UNSUPPORTED", f"源表 {sheet}!{neighbor_cell} 的金额单位无法按元/千元/万元解释")
            # A unit-bearing column header applies only to that amount column
            # (or a merged header spanning it), not to unrelated business columns.
            applies = col == target_col
            if reader.workbook and row <= target_row:
                for merged in reader.workbook[sheet].merged_cells.ranges:
                    if cell in merged and merged.min_col <= target_col <= merged.max_col:
                        applies = True
                        break
            if applies and row < target_row:
                currency_header = re.search(r"\((?:千|万)?(" + CUR_PATTERN + r")\)", value, re.I)
                if currency_header:
                    currency(cell, raw, currency_header[1])
                matched = re.search(r"\((?:人民币)?(万元|千元|元)\)", value)
                if matched:
                    unit(cell, raw, matched[1], "amount_column_header")
                elif value in UNITS:
                    unit(cell, raw, value, "amount_column_header")
        need(not any(item["currency"] != "CNY" for item in currencies), "FACT_FOREIGN_CURRENCY", "事实来源存在明确外币元数据，当前链路不自动换汇")
        if kind == "money":
            need(not unknown_units, "FACT_UNIT_UNSUPPORTED", "源表出现不能按元/千元/万元解释的金额单位声明")
            conflicts = [item for item in units if item["unit"] != declared_unit]
            need(not conflicts, "FACT_UNIT_CONFLICT", f"来源可见单位与 source.amount_unit={declared_unit} 不一致，不能忽略单位抄数")
        return units, currencies

    def resolve_node(self, node, fact_path):
        spec = node.get("source")
        need(isinstance(spec, dict), "FACT_SOURCE_SCHEMA", f"{fact_path}.source 必须为对象")
        kind = spec.get("kind", "money")
        need(kind in ("money", "number", "text"), "FACT_SOURCE_KIND", f"{fact_path} 的 kind 只能为 money/number/text")
        unit = spec.get("amount_unit")
        if kind == "money":
            need(unit in UNITS, "FACT_UNIT_REQUIRED", f"{fact_path} 必须显式声明 source.amount_unit（元/千元/万元），不能默认元")
        path, reader, actual_hash = self.reader(spec)
        sheet = spec.get("sheet")
        need(isinstance(sheet, str) and sheet, "FACT_SOURCE_SHEET", "必须指定事实来源工作表；CSV 使用 CSV")
        cell = cell_address(spec.get("cell"))
        raw = reader.value(sheet, cell)
        units, currencies = self.metadata(path, reader, sheet, cell, kind, unit)
        anchors = []
        if "anchor" in spec:
            anchor = spec["anchor"]
            need(isinstance(anchor, dict) and isinstance(anchor.get("text"), str) and anchor["text"], "FACT_ANCHOR_SCHEMA", "source.anchor 需要明确 cell/text")
            anchor_cell = cell_address(anchor.get("cell"))
            observed = reader.value(sheet, anchor_cell)
            need(observed == anchor["text"], "FACT_ANCHOR_MISMATCH", f"事实标签锚点不匹配: {path.name}!{sheet}!{anchor_cell}")
            anchors.append({"type": "label", "cell": anchor_cell, "expected": anchor["text"], "observed": json_value(observed)})
        identities = spec.get("identity_anchors", [])
        need(isinstance(identities, list), "FACT_IDENTITY_ANCHOR_SCHEMA", "identity_anchors 必须为列表")
        for anchor in identities:
            need(isinstance(anchor, dict) and anchor.get("context_key") in IDENTITY_KEYS, "FACT_IDENTITY_ANCHOR_SCHEMA", "身份锚点只支持 taxpayer_id/taxpayer_name/period_end")
            context_key = anchor["context_key"]
            need(context_key in self.context, "FACT_CONTEXT_MISSING", f"context 缺少身份锚点所需的 {context_key}")
            anchor_cell = cell_address(anchor.get("cell"))
            observed = reader.value(sheet, anchor_cell)
            if context_key == "period_end":
                actual = iso_date(observed, "事实来源期末日期锚点").isoformat()
                expected = iso_date(self.context[context_key], "context.period_end").isoformat()
            else:
                need(isinstance(observed, str) and isinstance(self.context[context_key], str), "FACT_IDENTITY_TEXT", "税号与企业名称须为文本，不能把数字型税号当成已核实")
                actual, expected = observed, self.context[context_key]
            need(actual == expected, "FACT_IDENTITY_MISMATCH", f"事实来源身份/期间锚点与 {context_key} 冲突")
            anchors.append({"type": "identity", "cell": anchor_cell, "context_key": context_key, "expected": expected, "observed": json_value(observed)})
        conversions = []
        if kind in ("money", "number"):
            numeric, conversions = parse_money(raw)
            value = Decimal(numeric)
            multiplier = UNITS[unit] if kind == "money" else Decimal(1)
            with localcontext() as precision:
                precision.prec = 50
                value *= multiplier
            normalized = format(value, "f")
            if "value" in node:
                original, _ = parse_money(node["value"])
                need(Decimal(original) == value, "FACT_VALUE_MISMATCH", f"{fact_path}.value 与源格归一化读值不符，拒绝悄悄覆盖抄错金额")
        else:
            need(isinstance(raw, (str, date, datetime)), "FACT_TEXT_TYPE", "text 来源须为真实文本或原生日期，数字身份可能已丢精度")
            normalized = str(serial(raw))
            need(not normalized.lstrip().startswith(("=", "+", "-", "@")), "FACT_TEXT_FORMULA", "text 来源具有公式前缀，不能作为普通文字取数")
            multiplier = None
            if "value" in node:
                need(isinstance(node["value"], str) and node["value"] == normalized, "FACT_VALUE_MISMATCH", f"{fact_path}.value 与来源原文不严格一致")
        binding = {"verification": "machine_read_source", "path": str(path), "sha256": actual_hash,
                   "sheet": sheet, "cell": cell, "kind": kind, "raw_value": json_value(raw), "normalized_value": normalized,
                   "source_amount_unit": unit if kind == "money" else None, "output_amount_unit": "元" if kind == "money" else None,
                   "multiplier": format(multiplier, "f") if multiplier is not None else None,
                   "unit_evidence_type": "visible_unit_and_source_declaration" if kind == "money" and units else "unit_declaration_only" if kind == "money" else "not_applicable",
                   "unit_observations": units, "currency_observations": currencies, "conversions": conversions,
                   "anchors": anchors, "context_keys_machine_checked": sorted({a["context_key"] for a in anchors if a["type"] == "identity"})}
        node["value"] = normalized
        if not isinstance(node.get("evidence"), str) or not node["evidence"].strip():
            node["evidence"] = f"实际读取 {path.name}!{sheet}!{cell}（SHA256 {actual_hash}）"
        node["source_binding"] = binding
        self.trace.append({"fact_path": fact_path, "evidence_type": "machine_read_source", "file_verified": True, "source_binding": deepcopy(binding)})

    def walk(self, value, fact_path="facts", depth=0):
        need(depth <= 40, "FACT_NESTING_LIMIT", "事实对象嵌套过深")
        if isinstance(value, dict):
            if "source" in value:
                self.resolve_node(value, fact_path)
            elif "value" in value and "evidence" in value:
                value["source_binding"] = {"verification": "explicit_declaration", "file_verified": False}
                self.trace.append({"fact_path": fact_path, "evidence_type": "explicit_declaration", "file_verified": False,
                                   "value": json_value(value["value"]), "evidence": value["evidence"]})
            for key, child in list(value.items()):
                if key not in ("source", "source_binding"):
                    self.walk(child, f"{fact_path}.{key}", depth + 1)
        elif isinstance(value, list):
            for index, child in enumerate(value):
                self.walk(child, f"{fact_path}[{index}]", depth + 1)

    def close(self):
        for reader in self.readers.values():
            if reader.workbook:
                reader.workbook.close()


def resolve_sources(facts, input_dir, context):
    """Return (new_facts, unique_file_manifests, per_fact_evidence_trace).

    Source-less value/evidence nodes remain explicit declarations. A genuine
    cell read proves that cell value, not taxpayer authenticity or human consent.
    Any source mismatch raises tax_fill.Blocked; no partial successful result is
    returned and the caller's object and input files remain unchanged.
    """
    need(isinstance(facts, dict), "FACT_SCHEMA", "facts/profile 必须为对象")
    result = deepcopy(facts)
    resolver = Resolver(input_dir, context)
    try:
        resolver.walk(result)
        for path, manifest in resolver.manifests.items():
            need(digest(path) == manifest["sha256"], "HASH_CHANGED", "事实来源在处理期间发生变化，必须重读后重新计算")
        return result, list(resolver.manifests.values()), resolver.trace
    finally:
        resolver.close()
